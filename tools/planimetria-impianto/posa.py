"""
Posa automatica dei simboli da tabella (Fase 4b).

Legge la tabella locali compilata (<PROGETTO>_locali.xlsx, colonne azzurre
con le quantità) e posa i simboli nei locali della tavola (contorni sul
layer LOCALI creati da aggiungi_locali.py):

  soffitto : distribuiti uniformemente nel locale (punti luce)
  parete   : distribuiti lungo il perimetro, staccati dal muro (prese, dati, TV)
  porta    : accanto alla porta, lato maniglia (comandi) / sopra la porta
             (lampade di emergenza). Le porte sono riconosciute dagli archi
             di apertura della pianta; senza porta -> come "parete".

Le posizioni sono una PROPOSTA da sistemare in AutoCAD LT. I simboli posati
sono marcati (xdata) e una nuova posa li sostituisce, senza duplicarli.

Uso:
    python posa.py "<tavola.dxf>" --tabella "<P>_locali.xlsx" --pianta "<pianta.dxf>" [--impianti LC,EM]
"""

import argparse
import json
import logging
import math
import re
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
import ezdxf
from ezdxf import xref

QUI = Path(__file__).resolve().parent
sys.path.insert(0, str(QUI))
logging.getLogger("ezdxf").setLevel(logging.ERROR)

import aggiungi_locali as al   # noqa: E402

APPID = "PLANIMETRIA_IMPIANTO"
LIBRERIA_DXF = QUI / "riferimenti" / "libreria_simboli.dxf"
LIBRERIA_JSON = QUI / "riferimenti" / "libreria_simboli.json"
LAYER_PORTE = ("Interni F", "PR_PORTE")

# modo di posa per blocco (default: parete)
MODO = {"CVE": "soffitto", "INT": "porta", "DEV": "porta", "P": "porta",
        "emer": "sopra_porta", "16A": "parete", "presa ethernet": "parete", "TV": "parete",
        "Rivelatore fumo in ambiente": "soffitto", "Riv Incendio Controsoffitto": "soffitto",
        "RIVI-P-01": "porta", "RIVS-SOA-01": "parete", "RIVA-CR-01_BIS": "parete",
        "CACC-SBA-01": "porta", "Controllo accessi": "porta", "Centralina accessi": "parete",
        "BLBAA": "parete", "Rack di piano": "parete", "Quadro di piano": "parete"}

STACCO_MURO = 25.0        # cm dal filo muro per i simboli a parete
MARGINE_SOFFITTO = 60.0   # cm minimi dal muro per i simboli a soffitto
DIST_MIN = 45.0           # cm minimi tra simboli
RAGGIO_PORTA = (60.0, 130.0)   # raggi degli archi considerati porte (cm)


# ── tabella ───────────────────────────────────────────────────────────────

def leggi_tabella(xlsx):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "PIANO")
    colonne = {}
    for j, v in enumerate(rows[h]):
        m = re.search(r"\[(.+)\]", str(v or ""))
        if m:
            colonne[j] = m.group(1)          # "blocco" oppure "blocco|LAYER"
    locali = []
    for r in rows[h + 1:]:
        if not r or not r[1]:
            continue
        q = {}
        for j, blocco in colonne.items():
            v = r[j]
            if v in (None, ""):
                continue
            n = float(v)
            if n < 0 or n != int(n):
                raise ValueError(f"Quantita' non valida '{v}' per {r[1]} / {blocco}")
            if n:
                q[blocco] = int(n)
        locali.append({"piano": r[0], "nome": r[1], "quantita": q})
    return locali


# ── geometria ─────────────────────────────────────────────────────────────

def _contiene(poly, pts):
    from matplotlib.path import Path as MPath
    return MPath(poly).contains_points(pts)


def _dist_bordo(poly, p):
    """Distanza minima di p dal contorno e punto più vicino."""
    best, bp, bi = 1e18, None, -1
    n = len(poly)
    for i in range(n):
        a, b = np.array(poly[i]), np.array(poly[(i + 1) % n])
        ab = b - a
        t = 0.0 if not ab.any() else max(0.0, min(1.0, np.dot(p - a, ab) / np.dot(ab, ab)))
        q = a + t * ab
        d = np.hypot(*(p - q))
        if d < best:
            best, bp, bi = d, q, i
    return best, bp, bi


def _normale_interna(poly, i, p_su_lato):
    a, b = np.array(poly[i]), np.array(poly[(i + 1) % len(poly)])
    d = (b - a) / (np.hypot(*(b - a)) or 1)
    n = np.array([-d[1], d[0]])
    prova = p_su_lato + n * 5
    return n if _contiene(poly, [prova])[0] else -n


def punti_soffitto(poly, n, margine=MARGINE_SOFFITTO):
    """n punti distribuiti uniformemente (k-means su una griglia interna)."""
    P = np.array(poly)
    x0, y0 = P.min(axis=0)
    x1, y1 = P.max(axis=0)
    passo = max(10.0, min(x1 - x0, y1 - y0) / 40)
    gx, gy = np.meshgrid(np.arange(x0, x1, passo), np.arange(y0, y1, passo))
    cand = np.c_[gx.ravel(), gy.ravel()]
    for m in (margine, margine / 2, 0):
        dentro = _contiene(poly, cand)
        if m:
            for dx, dy in ((m, 0), (-m, 0), (0, m), (0, -m)):
                dentro &= _contiene(poly, cand + [dx, dy])
        if dentro.sum() >= n:
            break
    C = cand[dentro]
    if len(C) == 0:
        return []
    rng = np.random.default_rng(0)
    centri = C[rng.choice(len(C), size=min(n, len(C)), replace=False)]
    for _ in range(25):
        lab = np.argmin(((C[:, None, :] - centri[None]) ** 2).sum(-1), axis=1)
        centri = np.array([C[lab == k].mean(axis=0) if (lab == k).any() else centri[k]
                           for k in range(len(centri))])
    # riporta ogni centro sul candidato interno più vicino
    return [tuple(C[np.argmin(((C - c) ** 2).sum(-1))]) for c in centri]


def punti_parete(poly, n, occupati, lato_min=80.0, dist_min=DIST_MIN, passo=20.0):
    """
    n punti lungo il perimetro (lati >= lato_min), staccati dal muro.
    Posizioni ideali equidistanti; se occupate si scorre lungo tutto il perimetro
    a passi di `passo`. Se non c'e' posto con dist_min, si riprova con meta' distanza.
    """
    lati = []
    for i in range(len(poly)):
        a, b = np.array(poly[i]), np.array(poly[(i + 1) % len(poly)])
        L = np.hypot(*(b - a))
        if L >= lato_min:
            lati.append((i, a, b, L))
    if not lati:
        lati = [(i, np.array(poly[i]), np.array(poly[(i + 1) % len(poly)]),
                 np.hypot(*(np.array(poly[(i + 1) % len(poly)]) - np.array(poly[i]))))
                for i in range(len(poly))]
    tot = sum(l[3] for l in lati)

    def punto(s):
        s %= tot
        for i, a, b, L in lati:
            if s <= L:
                margine = min(40.0, L / 4)
                t = max(margine, min(L - margine, s)) / L
                q = a + (b - a) * t
                nrm = _normale_interna(poly, i, q)
                p = q + nrm * STACCO_MURO
                return (p[0], p[1], _rotazione(nrm))
            s -= L
        return None

    out = []
    for dmin in (dist_min, dist_min / 2):
        for k in range(len(out), n):
            ideale = (k + 0.5) / n * tot
            for j in range(int(tot / passo) + 1):
                p = punto(ideale + j * passo)
                if p and all(np.hypot(p[0] - o[0], p[1] - o[1]) >= dmin for o in list(occupati) + out):
                    out.append(p)
                    break
        if len(out) == n:
            break
    return out


def _rotazione(normale_interna):
    """Rotazione (gradi) che porta la base del simbolo (lato -Y del blocco) verso il muro."""
    return (math.degrees(math.atan2(normale_interna[1], normale_interna[0])) - 90.0) % 360


def porte_vicine(poly, porte, toll=70.0):
    """Archi-porta il cui cardine è entro toll dal contorno del locale."""
    out = []
    for pr in porte:
        d, q, i = _dist_bordo(poly, np.array(pr["cardine"]))
        if d <= toll:
            out.append({**pr, "lato": i, "su_bordo": q})
    return out


def punto_porta(poly, porta, sopra=False):
    """Comando lato maniglia (o lampada sopra porta), staccato dal muro verso l'interno."""
    c = np.array(porta["cardine"])
    e1, e2 = np.array(porta["estremi"][0]), np.array(porta["estremi"][1])
    # anta chiusa = estremo più vicino al contorno del locale (giace sul filo muro)
    chiuso = e1 if _dist_bordo(poly, e1)[0] <= _dist_bordo(poly, e2)[0] else e2
    if sopra:
        bersaglio = c + (chiuso - c) / 2                       # centro dell'apertura
    else:
        v = (chiuso - c) / (np.hypot(*(chiuso - c)) or 1)
        bersaglio = chiuso + v * 25                            # oltre lo stipite lato maniglia
    d, q, i = _dist_bordo(poly, bersaglio)
    nrm = _normale_interna(poly, i, q)
    p = q + nrm * STACCO_MURO
    return (p[0], p[1], 0.0 if sopra else _rotazione(nrm))


# ── porte dalla pianta ────────────────────────────────────────────────────

def porte_in_tavola(pdoc, piante, f_unita=1.0):
    """Archi di apertura delle porte (coordinate tavola) dentro le zone di pianta."""
    out = []
    for e in pdoc.modelspace().query("ARC"):
        if e.dxf.layer not in LAYER_PORTE:
            continue
        r = e.dxf.radius * f_unita
        sweep = (e.dxf.end_angle - e.dxf.start_angle) % 360
        if not (RAGGIO_PORTA[0] <= r <= RAGGIO_PORTA[1]) or not (60 <= sweep <= 120):
            continue
        c = e.dxf.center
        for p in piante:
            x0, y0, x1, y1 = p["finestra"]
            if x0 <= c.x <= x1 and y0 <= c.y <= y1:
                s, (ix, iy) = p["scala"], p["ins"]
                T = lambda v: (ix + v[0] * s, iy + v[1] * s)
                out.append({"cardine": T((c.x, c.y)),
                            "estremi": [T((e.start_point.x, e.start_point.y)),
                                        T((e.end_point.x, e.end_point.y))],
                            "raggio": r * s})
    return out


# ── posa ──────────────────────────────────────────────────────────────────

def importa_blocco(tdoc, nome, lib):
    if nome in tdoc.blocks:
        return True
    if nome not in lib.blocks:
        return False
    loader = xref.Loader(lib, tdoc)
    loader.load_block_layout(lib.blocks.get(nome))
    loader.execute()
    return True


def locali_tavola(tdoc):
    out = []
    for pl in tdoc.modelspace().query('LWPOLYLINE[layer=="LOCALI"]'):
        try:
            xd = [v for _, v in pl.get_xdata(APPID)]
        except ezdxf.DXFValueError:
            continue
        if xd and xd[0] == "LOCALE":
            out.append({"nome": xd[1], "piano": xd[2], "poly": [tuple(p) for p in pl.get_points("xy")]})
    return out


def posa(tavola, tabella, pianta_dxf, impianti=None, out=None):
    tavola = Path(tavola)
    tdoc = ezdxf.readfile(tavola)
    msp = tdoc.modelspace()
    lib = ezdxf.readfile(LIBRERIA_DXF)
    info = json.loads(LIBRERIA_JSON.read_text(encoding="utf-8"))
    if APPID not in tdoc.appids:
        tdoc.appids.new(APPID)

    # rimuove la posa precedente
    for e in list(msp.query("INSERT")):
        try:
            if e.get_xdata(APPID)[0][1] == "POSA":
                msp.delete_entity(e)
        except ezdxf.DXFValueError:
            pass

    # blocchi di libreria non piu' usati in tavola: eliminati per reimportare
    # la versione aggiornata della libreria
    usati_in_tavola = {e.dxf.name for e in msp.query("INSERT")}
    for nome in [b.name for b in lib.blocks if not b.name.startswith("*")]:
        if nome in tdoc.blocks and nome not in usati_in_tavola:
            tdoc.blocks.delete_block(nome, safe=False)

    pdoc = ezdxf.readfile(pianta_dxf)
    piante = al.piante_in_tavola(tdoc, Path(pianta_dxf).stem)
    porte = porte_in_tavola(pdoc, piante)

    locali = locali_tavola(tdoc)
    per_chiave = defaultdict(list)
    for l in locali:
        per_chiave[(l["piano"], l["nome"])].append(l)
    usati = defaultdict(int)

    avvisi, posati = [], defaultdict(int)
    occupati_per_poly = defaultdict(list)
    for riga in leggi_tabella(tabella):
        if not riga["quantita"]:
            continue
        k = (riga["piano"], riga["nome"])
        cand = per_chiave.get(k, [])
        if usati[k] >= len(cand):
            avvisi.append(f"{riga['nome']} ({riga['piano']}): locale non trovato in tavola")
            continue
        loc = cand[usati[k]]
        usati[k] += 1
        poly = loc["poly"]
        chiave_poly = tuple(poly[0])
        occupati = occupati_per_poly[chiave_poly]
        porte_loc = porte_vicine(poly, porte)
        for voce, n in riga["quantita"].items():
            blocco, _, layer_col = voce.partition("|")
            meta = info.get(blocco)
            layer = layer_col or (meta or {}).get("layer") or "0"
            if layer not in tdoc.layers:
                tdoc.layers.add(layer)
            if impianti and layer not in impianti:
                continue
            if not importa_blocco(tdoc, blocco, lib):
                avvisi.append(f"blocco '{blocco}' non presente in libreria: non posato")
                continue
            modo = MODO.get(blocco, "parete")
            if modo == "soffitto":
                pts = [(x, y, 0.0) for x, y in punti_soffitto(poly, n)]
            elif modo in ("porta", "sopra_porta") and porte_loc:
                pts = [punto_porta(poly, porte_loc[i % len(porte_loc)], sopra=(modo == "sopra_porta"))
                       for i in range(n)]
                # piu' comandi sulla stessa porta: affiancati
                pts = [(p[0] + (i // len(porte_loc)) * DIST_MIN, p[1], p[2]) for i, p in enumerate(pts)]
            else:
                pts = punti_parete(poly, n, occupati)
            if len(pts) < n:
                avvisi.append(f"{riga['nome']}: posati {len(pts)} di {n} '{blocco}' (spazio insufficiente)")
            s = (meta or {}).get("scala", 1.0)
            for p in pts:
                ins = msp.add_blockref(blocco, p[:2], dxfattribs={"layer": layer, "xscale": s,
                                                                  "yscale": s, "zscale": s,
                                                                  "rotation": p[2]})
                if tdoc.blocks.get(blocco).query("ATTDEF"):
                    ins.add_auto_attribs({})
                ins.set_xdata(APPID, [(1000, "POSA"), (1000, riga["nome"]), (1000, str(riga["piano"]))])
                occupati.append(p[:2])
                posati[blocco] += 1

    out = Path(out) if out else tavola
    tdoc.saveas(out)
    return {"tavola": str(out), "posati": dict(posati), "porte_riconosciute": len(porte),
            "avvisi": avvisi}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("tavola")
    ap.add_argument("--tabella", required=True)
    ap.add_argument("--pianta", required=True)
    ap.add_argument("--impianti", help="layer da posare, es. LC,EM (default: tutti)")
    ap.add_argument("--out")
    a = ap.parse_args()
    imp = set(a.impianti.split(",")) if a.impianti else None
    print(json.dumps(posa(a.tavola, a.tabella, a.pianta, imp, a.out), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
