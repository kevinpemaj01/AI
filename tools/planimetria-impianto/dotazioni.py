"""
Dotazioni per locale da un file di regole (Fase 5).

Compila le quantità della tabella locali (<P>_locali.xlsx) applicando un
file di regole JSON (cartella regole/):

  - tipo del locale riconosciuto dal nome (parole chiave, in ordine)
  - quantità per tipo: {"fisso": n} oppure {"min": n, "ogni_m2": A}
    -> max(min, ceil(superficie / A))
  - "per_piano": un elemento per piano nel primo locale del tipo preferito
  - "per_edificio": un elemento per edificio
  - i locali segnalati "stesso spazio di X" sono già compresi in X: nessuna dotazione

Le celle già compilate dall'utente NON vengono sovrascritte (salvo --sovrascrivi).

Il file di regole è la fonte dei valori: lo script non contiene numeri di
progetto. regole/esempio_dimostrativo.json contiene valori DIMOSTRATIVI;
i valori CEI 64-8 vanno inseriti in un file di regole dedicato.

Uso:
    python dotazioni.py "<P>_locali.xlsx" --regole regole/<file>.json [--sovrascrivi]
"""

import argparse
import json
import math
import sys
from pathlib import Path

QUI = Path(__file__).resolve().parent
sys.path.insert(0, str(QUI))

import aggiungi_locali as al  # noqa: E402


def leggi_righe(xlsx: Path):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "PIANO")
    out = []
    for r in rows[h + 1:]:
        if not r or not r[1]:
            continue
        out.append({"piano": r[0], "nome": r[1], "superficie_calcolata": r[2],
                    "superficie_dichiarata": r[3], "scostamento_%": r[4], "esito": r[5], "note": r[6] or ""})
    return out


def tipo_locale(nome, regole):
    n = nome.upper()
    for t in regole["tipi"]:
        if any(p.upper() in n for p in t["parole"]):
            return t
    return None


def quantita(regola, superficie):
    if "fisso" in regola:
        return int(regola["fisso"])
    q = regola.get("min", 0)
    if regola.get("ogni_m2") and superficie:
        q = max(q, math.ceil(superficie / regola["ogni_m2"]))
    return int(q)


def intestazione(col):
    voce = col["blocco"] + (f"|{col['layer']}" if col.get("layer") else "")
    return f"{col['nome']}\n[{voce}]"


def calcola(righe, regole):
    col_per_blocco = {c["blocco"]: intestazione(c) for c in regole["colonne"]}
    avvisi = []
    for r in righe:
        r["quantita"] = {}
        r["tipo"] = None
        if "stesso spazio di" in (r["note"] or ""):
            continue                                   # compreso nel locale principale
        t = tipo_locale(r["nome"], regole)
        if t is None:
            avvisi.append(f"{r['nome']} ({r['piano']}): tipo di locale non riconosciuto, nessuna dotazione")
            continue
        r["tipo"] = t["tipo"]
        sup = r["superficie_calcolata"] or 0
        for blocco, reg in t["dotazioni"].items():
            n = quantita(reg, sup)
            if n:
                r["quantita"][col_per_blocco[blocco]] = n

    def assegna(blocco, preferenze, gruppo):
        for pref in preferenze:
            for r in gruppo:
                if r["tipo"] == pref:
                    k = col_per_blocco[blocco]
                    r["quantita"][k] = r["quantita"].get(k, 0) + 1
                    return True
        avvisi.append(f"'{blocco}': nessun locale adatto ({', '.join(preferenze)})")
        return False

    piani = []
    for r in righe:
        if r["piano"] not in piani:
            piani.append(r["piano"])
    for voce in regole.get("per_piano", []):
        for p in piani:
            assegna(voce["blocco"], voce["preferenze"], [r for r in righe if r["piano"] == p])
    for voce in regole.get("per_edificio", []):
        assegna(voce["blocco"], voce["preferenze"], righe)
    return avvisi


def compila(xlsx, file_regole, progetto=None, sovrascrivi=False):
    xlsx, file_regole = Path(xlsx), Path(file_regole)
    regole = json.loads(file_regole.read_text(encoding="utf-8"))
    righe = leggi_righe(xlsx)
    avvisi = calcola(righe, regole)
    if not sovrascrivi:                                # non toccare le celle gia' compilate
        _, esistenti = al.leggi_esistente(xlsx)
        occ = {}
        for r in righe:
            k = (r["piano"], r["nome"])
            occ[k] = occ.get(k, 0) + 1
            gia = esistenti.get((r["piano"], r["nome"], occ[k]), {})
            r["quantita"] = {c: v for c, v in r["quantita"].items() if c not in gia}
    # con sovrascrivi: i valori calcolati prendono il posto di quelli esistenti
    # (scrivi_tabella fa prevalere r["quantita"] sui valori gia' presenti)
    nota = f"{regole.get('etichetta', '')} - regole: {file_regole.name}. " \
           "Righe gialle: da verificare."
    colonne = [intestazione(c) for c in regole["colonne"]]
    al.scrivi_tabella(righe, xlsx, progetto or xlsx.stem.replace("_locali", ""), extra_colonne=colonne,
                      nota=nota)
    totali = {}
    for r in righe:
        for c, v in r["quantita"].items():
            totali[c.split("[")[1].rstrip("]")] = totali.get(c.split("[")[1].rstrip("]"), 0) + v
    return {"tabella": str(xlsx), "regole": file_regole.name, "totali_aggiunti": totali,
            "tipi": {f"{r['piano']} / {r['nome']}": r["tipo"] for r in righe}, "avvisi": avvisi}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("tabella")
    ap.add_argument("--regole", required=True)
    ap.add_argument("--progetto")
    ap.add_argument("--sovrascrivi", action="store_true")
    a = ap.parse_args()
    print(json.dumps(compila(a.tabella, a.regole, a.progetto, a.sovrascrivi), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
