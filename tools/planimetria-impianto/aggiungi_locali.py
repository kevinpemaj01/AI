"""
Aggiunge i locali a una tavola impianto e crea la tabella locali (Excel).

Per ogni pianta inserita nella tavola (xref ritagliato creato da
prepara_tavola.py) riconosce i locali con locali.py, poi:

  - disegna il contorno di ogni locale sul layer LOCALI della tavola
    (layer NON stampabile: serve come riferimento per posa e verifiche)
  - crea <PROGETTO>_locali.xlsx con: piano, locale, superfici, esito,
    note e le colonne delle dotazioni da compilare (input della posa)

Uso:
    python aggiungi_locali.py "<tavola.dxf>" --pianta "<pianta.dxf>" --progetto VJN001 [--unita cm]
"""

import argparse
import json
import logging
import sys
from pathlib import Path

import ezdxf
from ezdxf import xclip

QUI = Path(__file__).resolve().parent
sys.path.insert(0, str(QUI))
sys.path.insert(0, str(QUI.parent / "cad-output"))
logging.getLogger("ezdxf").setLevel(logging.ERROR)

import locali as lc          # noqa: E402
import dxf_writer as dw      # noqa: E402

# colonne dotazioni nella tabella (da compilare a mano; codice blocco del catalogo)
COLONNE_DOTAZIONI = [
    ("PUNTI LUCE", "CVE"),
    ("COMANDI (int./dev./puls.)", "INT"),
    ("PRESE 10/16A", "16A"),
    ("PRESE DATI", "presa ethernet"),
    ("PRESE TV", "TV"),
    ("LAMPADE EMERGENZA", "emer"),
]


def piante_in_tavola(tdoc, nome_pianta):
    """Inserimenti dell'xref della pianta: (finestra in coord. pianta, inserimento, scala)."""
    out = []
    for ins in tdoc.modelspace().query("INSERT"):
        if ins.dxf.name != nome_pianta:
            continue
        c = xclip.XClip(ins)
        if not c.has_clipping_path:
            continue
        v = c.get_block_clipping_path().vertices
        xs, ys = [p.x for p in v], [p.y for p in v]
        out.append({"finestra": (min(xs), min(ys), max(xs), max(ys)),
                    "ins": (ins.dxf.insert.x, ins.dxf.insert.y), "scala": ins.dxf.xscale})
    return out


def etichetta_piano(tdoc, p):
    """Testo sottolineato sotto la pianta (es. 'PIANTA PIANO TERRA ...')."""
    x0, y0, x1, y1 = p["finestra"]
    s = p["scala"]
    cx = p["ins"][0] + (x0 + x1) / 2 * s
    ybot = p["ins"][1] + y0 * s
    migliore, dmin = None, 1e18
    for t in tdoc.modelspace().query("TEXT"):
        if t.dxf.layer != "TESTI" or not t.dxf.text.startswith("%%U"):
            continue
        pt = t.dxf.align_point if t.dxf.halign else t.dxf.insert
        d = abs(pt.x - cx) + abs(pt.y - ybot)
        if d < dmin:
            migliore, dmin = t.dxf.text[3:], d
    return migliore or "PIANTA"


def aggiungi(tavola, pianta_dxf, progetto, unita="cm", out_dir=None):
    tavola, pianta_dxf = Path(tavola), Path(pianta_dxf)
    tdoc = ezdxf.readfile(tavola)
    pdoc = ezdxf.readfile(pianta_dxf)
    nome_pianta = pianta_dxf.stem
    piante = piante_in_tavola(tdoc, nome_pianta)
    if not piante:
        raise ValueError(f"Nella tavola non c'e' nessun xref ritagliato di '{nome_pianta}'.")

    if "PLANIMETRIA_IMPIANTO" not in tdoc.appids:
        tdoc.appids.new("PLANIMETRIA_IMPIANTO")
    if "LOCALI" not in tdoc.layers:
        tdoc.layers.add("LOCALI", color=8)
    tdoc.layers.get("LOCALI").dxf.plot = 0            # non stampabile
    msp = tdoc.modelspace()
    for e in list(msp.query('*[layer=="LOCALI"]')):   # rigenera i contorni
        msp.delete_entity(e)

    righe = []
    for p in piante:
        piano = etichetta_piano(tdoc, p)
        s, (ix, iy) = p["scala"], p["ins"]
        for r in lc.trova_locali(pdoc, p["finestra"], unita):
            if "poligono" in r and r["esito"] != "NON TROVATO":
                pts = [(ix + x * s, iy + y * s) for x, y in r["poligono"]]
                pl = msp.add_lwpolyline(pts, close=True, dxfattribs={"layer": "LOCALI"})
                pl.set_xdata("PLANIMETRIA_IMPIANTO", [(1000, "LOCALE"), (1000, r["nome"]),
                                                      (1000, piano)])
            righe.append({"piano": piano, **r})

    tdoc.saveas(tavola)
    out_dir = Path(out_dir) if out_dir else dw.project_dir(progetto)
    xlsx = scrivi_tabella(righe, out_dir / f"{progetto}_locali.xlsx", progetto)
    return {"tavola": str(tavola), "tabella": str(xlsx), "locali": righe,
            "da_verificare": sum(1 for r in righe if r["esito"] != "OK")}


N_BASE = 7   # colonne fisse: PIANO ... NOTE


def _intestazione_dotazione(nome, blocco):
    return f"{nome}\n[{blocco}]"


def leggi_esistente(xlsx: Path):
    """
    Tabella gia' esistente: (intestazioni dotazioni, {(piano, nome, n-esima occorrenza): {intestazione: valore}}).
    Serve a non perdere quantita' e colonne aggiunte (dall'utente o da dotazioni.py).
    """
    if not xlsx.exists():
        return [], {}
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    try:
        h = next(i for i, r in enumerate(rows) if r and r[0] == "PIANO")
    except StopIteration:
        return [], {}
    intest = [str(v) for v in rows[h][N_BASE:] if v]
    occ, out = {}, {}
    for r in rows[h + 1:]:
        if not r or not r[1]:
            continue
        k = (r[0], r[1])
        occ[k] = occ.get(k, 0) + 1
        vals = {intest[j]: v for j, v in enumerate(r[N_BASE:N_BASE + len(intest)]) if v not in (None, "")}
        if vals:
            out[(r[0], r[1], occ[k])] = vals
    return intest, out


def quantita_esistenti(xlsx: Path):
    """Compatibilita': {(piano, nome, occorrenza): {intestazione: valore}}."""
    return leggi_esistente(xlsx)[1]


def scrivi_tabella(righe, out: Path, progetto, extra_colonne=(), nota=None):
    """
    righe         : locali (da trova_locali) con "piano"
    extra_colonne : intestazioni dotazione aggiuntive "NOME\n[blocco]" o "NOME\n[blocco|LAYER]"
    I valori gia' presenti nella tabella esistente vengono conservati.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    vecchie_int, vecchie = leggi_esistente(out)
    dot = [_intestazione_dotazione(n, b) for n, b in COLONNE_DOTAZIONI]
    for c in list(vecchie_int) + list(extra_colonne):
        if c not in dot:
            dot.append(c)

    wb = Workbook()
    ws = wb.active
    ws.title = "Locali"
    ws.append([f"TABELLA LOCALI - {progetto}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([nota or "Superfici calcolate dai muri della pianta. Righe gialle: da verificare. "
               "Compilare le colonne azzurre con le quantita' da posare."])
    ws.append([])
    base = ["PIANO", "LOCALE", "SUP. CALCOLATA m²", "SUP. IN PIANTA m²", "SCOSTAMENTO %", "ESITO", "NOTE"]
    ws.append(base + dot)
    hdr = ws.max_row
    grigio, azzurro = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="DDEBF7")
    giallo = PatternFill("solid", fgColor="FFF2CC")
    for c in range(1, len(base) + len(dot) + 1):
        cell = ws.cell(row=hdr, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = azzurro if c > len(base) else grigio
    occ = {}
    for r in righe:
        k = (r["piano"], r["nome"])
        occ[k] = occ.get(k, 0) + 1
        q = vecchie.pop((r["piano"], r["nome"], occ[k]), {})
        q.update(r.get("quantita", {}))
        ws.append([r["piano"], r["nome"], r.get("superficie_calcolata"), r.get("superficie_dichiarata"),
                   r.get("scostamento_%"), r.get("esito"), r.get("note")] + [q.get(c) for c in dot])
        if r.get("esito") != "OK":
            for c in range(1, len(base) + 1):
                ws.cell(row=ws.max_row, column=c).fill = giallo
        for c in range(len(base) + 1, len(base) + len(dot) + 1):
            ws.cell(row=ws.max_row, column=c).fill = azzurro
    for i, w in enumerate([34, 26, 12, 12, 12, 14, 55] + [14] * len(dot), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hdr].height = 60
    ws.freeze_panes = ws.cell(row=hdr + 1, column=3)
    if vecchie:                                      # locali spariti: quantita' non perse
        ws2 = wb.create_sheet("Quantita' non abbinate")
        ws2.append(["PIANO", "LOCALE", "OCCORRENZA"] + dot)
        for (piano, nome, k), q in vecchie.items():
            ws2.append([piano, nome, k] + [q.get(c) for c in dot])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("tavola")
    ap.add_argument("--pianta", required=True)
    ap.add_argument("--progetto", required=True)
    ap.add_argument("--unita", default="cm")
    a = ap.parse_args()
    res = aggiungi(a.tavola, a.pianta, a.progetto, a.unita)
    print(json.dumps({k: v for k, v in res.items() if k != "locali"}, ensure_ascii=False, indent=2))
    for r in res["locali"]:
        print(f"  {r['esito']:13} {r['piano'][:28]:28} {r['nome'][:30]:30} "
              f"{r.get('superficie_calcolata', '-')!s:>6} m²  {r['note']}")


if __name__ == "__main__":
    main()
