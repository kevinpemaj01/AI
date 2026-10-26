"""
LOCK test runner for tools/planimetria-impianto

TP1 rebuilds sheet E01 with prepara_tavola.py and compares frame, cartiglio
xref, plan xref insertion and clip boundary against the user's real
ALLEGATO E01 (all coordinates relative to the outer frame corner).

Usage:
    python tools/planimetria-impianto/tests/run_lock.py
"""

import json
import logging
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
logging.getLogger("ezdxf").setLevel(logging.ERROR)

import ezdxf                      # noqa: E402
from ezdxf import bbox, xclip     # noqa: E402
import prepara_tavola as pt       # noqa: E402
import locali as lc               # noqa: E402
import aggiungi_locali as al      # noqa: E402
import posa as ps                 # noqa: E402
import computo as cp              # noqa: E402
import dotazioni as dt            # noqa: E402
import legenda as lg              # noqa: E402

EXP = json.loads((HERE / "expected_results.json").read_text(encoding="utf-8"))
results = []


def record(case, check, ok, got=None, want=None):
    results.append({"case": case, "check": check, "pass": bool(ok), "got": got, "want": want})


def misure_tavola(path):
    """Geometria chiave della tavola, relativa all'angolo della cornice esterna."""
    doc = ezdxf.readfile(path)
    msp = doc.modelspace()
    chiuse = [p for p in msp.query("LWPOLYLINE") if p.closed]
    cornice = max(chiuse, key=lambda p: bbox.extents([p]).size.x * bbox.extents([p]).size.y)
    e = bbox.extents([cornice])
    ox, oy = e.extmin.x, e.extmin.y
    out = {"cornice": [e.size.x, e.size.y], "xref": {}}
    for ins in msp.query("INSERT"):
        blk = doc.blocks.get(ins.dxf.name)
        if blk is None or not blk.block.dxf.flags & 4:   # (blocco xref vuoto = falsy)
            continue
        d = {"ins": [ins.dxf.insert.x - ox, ins.dxf.insert.y - oy], "scala": ins.dxf.xscale}
        c = xclip.XClip(ins)
        if c.has_clipping_path:
            v = c.get_wcs_clipping_path().vertices
            xs, ys = [p.x - ox for p in v], [p.y - oy for p in v]
            d["clip"] = [min(xs), min(ys), max(xs), max(ys)]
        out["xref"][ins.dxf.name] = d
    return doc, out


def vicino(a, b, tol):
    return all(abs(x - y) <= tol for x, y in zip(a, b))


def tp1(tmp):
    case = "TP1_ricostruisci_E01"
    c = EXP["prepara_tavola"][case]
    i, e, tol = c["inputs"], c["expected"], c["tolleranza"]
    res = pt.prepara("LOCKTEST", i["numero"], i["titolo"], ROOT / i["pianta"], i["unita"],
                     tuple(i["finestra"]), scala=i["scala"], out_dir=tmp)
    doc, got = misure_tavola(res["tavola"])
    _, ref = misure_tavola(ROOT / c["riferimento_E01"])

    record(case, "unita stimata", res["pianta"]["unita_stimata"] == e["unita_stimata"],
           res["pianta"]["unita_stimata"], e["unita_stimata"])
    record(case, "cornice", vicino(got["cornice"], ref["cornice"], tol), got["cornice"], ref["cornice"])
    for nome in ("Planimetria di Base", "CARTIGLIO BASE"):
        g, r = got["xref"].get(nome, {}), ref["xref"].get(nome, {})
        record(case, f"{nome}: inserimento", g and vicino(g["ins"], r["ins"], tol), g.get("ins"), r.get("ins"))
        record(case, f"{nome}: scala", g and abs(g["scala"] - r["scala"]) < 1e-9, g.get("scala"), r.get("scala"))
    g, r = got["xref"]["Planimetria di Base"], ref["xref"]["Planimetria di Base"]
    record(case, "ritaglio pianta", "clip" in g and vicino(g["clip"], r["clip"], tol), g.get("clip"), r.get("clip"))

    paths = {b.name: b.block.dxf.xref_path for b in doc.blocks if b.block.dxf.flags & 4}
    record(case, "percorsi xref", paths == e["xref"], paths, e["xref"])
    copiati = sorted(p.name for p in Path(tmp).glob("*.dwg"))
    record(case, "dwg copiati", copiati == sorted(e["xref_files_copied"]), copiati, e["xref_files_copied"])
    lay = {n: doc.layers.get(n).dxf.color for n in e["layers_impianti"] if n in doc.layers}
    record(case, "layer impianti", lay == e["layers_impianti"], lay, e["layers_impianti"])
    grigi = [l.dxf.color for l in doc.layers if l.dxf.name.startswith("Planimetria di Base|")]
    record(case, "layer pianta in grigio", grigi and all(c == e["pianta_layer_color"] for c in grigi),
           sorted(set(grigi)), e["pianta_layer_color"])
    hdr = {k: doc.header.get(k) for k in e["header"]}
    record(case, "header", hdr == e["header"], hdr, e["header"])


def tp2(tmp):
    case = "TP2_pianta_troppo_grande"
    c = EXP["prepara_tavola"][case]
    i, e = c["inputs"], c["expected"]
    try:
        pt.prepara("LOCKTEST", "E99", "PROVA", ROOT / i["pianta"], i["unita"], i["finestra"],
                   scala=i["scala"], out_dir=tmp)
        record(case, "errore atteso", False, "nessun errore", e["errore_contiene"])
    except ValueError as ex:
        record(case, "errore atteso", e["errore_contiene"] in str(ex), str(ex), e["errore_contiene"])


def _locali(case):
    c = EXP["locali"]
    t = c[case]
    doc = ezdxf.readfile(ROOT / c["pianta"])
    res = {r["nome"]: r for r in lc.trova_locali(doc, tuple(t["finestra"]), "cm")}
    tol = c["tolleranza_%"]
    for nome, sup in t["superfici"].items():
        r = res.get(nome)
        got = r.get("superficie_calcolata") if r else None
        ok = got is not None and abs(got - sup) / sup * 100 <= tol and r["esito"] in ("OK", "DA VERIFICARE")
        record(case, f"superficie {nome}", ok, got, f"{sup} ±{tol}%")
    for a, b in t["uniti"]:
        ra, rb = res.get(a), res.get(b)
        ok = bool(ra and rb and "stesso spazio" in (ra["note"] + rb["note"]))
        record(case, f"segnalati uniti {a} + {b}", ok, (ra or {}).get("note"), "stesso spazio")


def tl1(_tmp):
    _locali("TL1_stato_di_fatto_PT")


def tl2(_tmp):
    _locali("TL2_stato_di_fatto_P1")


def tpo1(tmp):
    from collections import Counter
    from matplotlib.path import Path as MPath
    case = "TPO1_posa_da_tabella"
    c = EXP["posa"]
    fin = [{"finestra": tuple(w), "etichetta": e} for w, e in zip(c["finestre"], c["etichette"])]
    pianta = ROOT / c["pianta"]
    for num, impianti in c["tavole"].items():
        r = pt.prepara("LOCKTEST", num, f"PROVA {num}", pianta, "cm", finestre=fin, out_dir=tmp)
        al.aggiungi(r["tavola"], pianta.with_suffix(".dxf"), "LOCKTEST", out_dir=tmp)
        for giro in (1, 2):                              # il 2o giro verifica: niente duplicati
            res = ps.posa(r["tavola"], ROOT / c["tabella"], pianta.with_suffix(".dxf"), set(impianti))
        doc = ezdxf.readfile(r["tavola"])
        loc = ps.locali_tavola(doc)
        n, fuori = Counter(), 0
        for e in doc.modelspace().query("INSERT"):
            try:
                xd = [v for _, v in e.get_xdata(ps.APPID)]
            except ezdxf.DXFValueError:
                continue
            n[e.dxf.name] += 1
            polys = [l["poly"] for l in loc if l["nome"] == xd[1] and l["piano"] == xd[2]]
            if not any(MPath(pl).contains_point((e.dxf.insert.x, e.dxf.insert.y)) for pl in polys):
                fuori += 1
        att = c["attesi"][num]
        record(case, f"{num}: quantita' posate = tabella (dopo 2 pose)", dict(n) == att, dict(n), att)
        record(case, f"{num}: simboli dentro il proprio locale", fuori == 0, fuori, 0)
        record(case, f"{num}: nessun avviso di posa", not res["avvisi"], res["avvisi"], [])
        record(case, f"{num}: audit DXF senza errori", len(doc.audit().errors) == 0, len(doc.audit().errors), 0)
        comp = {x["blocco"]: x["quantita"] for x in cp.computo(Path(r["tavola"]))["righe"] if x["blocco"] in att}
        record(case, f"{num}: computo = posa", comp == att, comp, att)


def td1(tmp):
    import shutil
    from openpyxl import load_workbook
    case = "TD1_dotazioni_da_regole"
    c = EXP["dotazioni"]
    x = Path(tmp) / "TEST_locali.xlsx"
    shutil.copy(ROOT / c["tabella"], x)
    wb = load_workbook(x)
    ws = wb.active
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "PIANO")
    for r in range(hdr + 1, ws.max_row + 1):            # tabella vuota...
        for col in range(8, ws.max_column + 1):
            ws.cell(r, col).value = None
    pre = c["cella_preimpostata"]                         # ...tranne una cella dell'utente
    col_pre = next(col for col in range(8, ws.max_column + 1) if str(ws.cell(hdr, col).value).startswith(pre["colonna"]))
    riga_pre = next(r for r in range(hdr + 1, ws.max_row + 1) if ws.cell(r, 2).value == pre["locale"])
    ws.cell(riga_pre, col_pre).value = pre["valore"]
    wb.save(x)
    dt.compila(x, ROOT / c["regole"], "TEST")
    righe, tutte = {}, []
    wb = load_workbook(x, data_only=True)
    ws = wb.active
    intest = {col: str(ws.cell(hdr, col).value) for col in range(8, ws.max_column + 1)}
    for r in range(hdr + 1, ws.max_row + 1):
        q = {}
        for col, h in intest.items():
            v = ws.cell(r, col).value
            if v not in (None, ""):
                q[h.split("[")[1].rstrip("]").split("|")[0]] = v
        righe.setdefault(ws.cell(r, 2).value, q)
        tutte.append(q)
    for locale, att in c["attesi"].items():
        got = {k: righe[locale].get(k) for k in att} if att else righe[locale]
        record(case, f"{locale}", got == att, got, att)
    badge = sum(q.get("CACC-SBA-01", 0) for q in tutte)      # anche i nomi ripetuti
    record(case, "lettori badge solo camere B&B", badge == c["totale_lettori_badge"], badge, c["totale_lettori_badge"])
    got = righe[pre["locale"]].get("CVE")
    record(case, "cella dell'utente non sovrascritta", got == pre["valore"], got, pre["valore"])


def tlg1(tmp):
    from collections import Counter
    case = "TLG1_legenda"
    c = EXP["posa"]
    fin = [{"finestra": tuple(w), "etichetta": e} for w, e in zip(c["finestre"], c["etichette"])]
    pianta = ROOT / c["pianta"]
    r = pt.prepara("LOCKTEST", "E02", "PROVA LEGENDA", pianta, "cm", finestre=fin, out_dir=tmp)
    al.aggiungi(r["tavola"], pianta.with_suffix(".dxf"), "LOCKTEST", out_dir=tmp)
    ps.posa(r["tavola"], ROOT / c["tabella"], pianta.with_suffix(".dxf"), set(c["tavole"]["E02"]))
    prima = cp.computo(Path(r["tavola"]))["totale_simboli"]
    for _ in (1, 2):                                      # rigenerata due volte: niente duplicati
        g = lg.genera(r["tavola"], ROOT / EXP["dotazioni"]["regole"])
    att = len(c["attesi"]["E02"])
    record(case, "una voce per simbolo", g["voci_legenda"] == att, g["voci_legenda"], att)
    record(case, "tutte le voci con descrizione", not g["senza_descrizione"], g["senza_descrizione"], [])
    dopo = cp.computo(Path(r["tavola"]))["totale_simboli"]
    record(case, "la legenda non cambia il computo", dopo == prima, dopo, prima)
    doc = ezdxf.readfile(r["tavola"])
    n_leg = sum(1 for e in doc.modelspace().query("INSERT")
                if e.has_xdata(ps.APPID) and e.get_xdata(ps.APPID)[0][1] == "LEGENDA")
    record(case, "simboli di legenda non duplicati", n_leg == att, n_leg, att)


def main():
    for fn in (tp1, tp2, tl1, tl2, tpo1, td1, tlg1):
        with tempfile.TemporaryDirectory() as tmp:
            try:
                fn(tmp)
            except Exception as ex:
                record(fn.__name__, "exception", False, repr(ex), None)
    failed = [r for r in results if not r["pass"]]
    for r in results:
        line = f"[{'PASS' if r['pass'] else 'FAIL'}] {r['case']}: {r['check']}"
        if not r["pass"]:
            line += f"  got={r['got']}  want={r['want']}"
        print(line)
    print(f"\n{len(results) - len(failed)}/{len(results)} checks passed.")
    print("LOCK: PASS" if not failed else "LOCK: FAIL")
    sys.exit(0 if not failed else 1)


if __name__ == "__main__":
    main()
