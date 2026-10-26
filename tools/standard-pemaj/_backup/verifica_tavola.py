"""
Studio Pemaj - Controllo di una tavola DXF contro il prontuario di disegno.

Uso:
  python verifica_tavola.py "<tavola.dxf>" [--unita cm|mm|m] [--scala 100] [--tipo progetto|diri|schema]
                            [--progetto <P>]

  --unita / --scala servono per calcolare l'altezza dei testi sulla carta (default: cm e 1:100,
  la convenzione delle tavole dello studio). I testi con stile annotativo non vengono misurati.

Esce con un JSON (status ok / errori) e scrive il report Excel:
  Projects/<P>/<tavola>_verifica.xlsx   se --progetto e' indicato
  out/verifiche/<tavola>_verifica.xlsx  altrimenti

Codici: V01 layer non standard, V02 colore forzato, V03 spessore forzato, V04 testo capovolto,
V05 testo sotto 1,8 mm su carta, V06 stile di stampa non standard, V07 cartiglio,
V08 scala del cartiglio diversa dalla finestra, V09 legenda DI.RI. senza "esistente".
"""
import argparse
import sys
from collections import Counter, defaultdict
from pathlib import Path

import ezdxf

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comune import (LAYERS, LAYERS_SERVIZIO, ALTEZZA_MIN_MM, CTB_AMMESSI, OUT, BASE,  # noqa: E402
                    esci)

MM_PER_UNITA = {"mm": 1.0, "cm": 10.0, "m": 1000.0}
STANDARD = {n.upper() for n, _, _ in LAYERS} | LAYERS_SERVIZIO
SEGNAPOSTO = {"ALLEGATO E00", "TITOLO DELLA TAVOLA"}
MAX_DETTAGLI = 200


def stile_annotativo(doc, nome):
    st = doc.styles.get(nome) if nome in doc.styles else None
    return bool(st is not None and st.has_xdata("AcadAnnotative"))


def verifica(path, unita="cm", scala=100.0, tipo="progetto"):
    doc = ezdxf.readfile(str(path))
    msp = doc.modelspace()
    trovati = []                                   # (codice, gravita', regola, dettaglio)

    def add(cod, grav, regola, dett):
        trovati.append((cod, grav, regola, dett))

    # V01 - layer non standard (esclusi xref "X|Y")
    uso = Counter(e.dxf.layer for e in msp if e.dxf.hasattr("layer"))
    for lay, n in sorted(uso.items()):
        u = lay.upper()
        if "|" in u or u in STANDARD:
            continue
        nota = " (importato da PDF)" if u.startswith("PDF") else ""
        add("V01", "AVVISO", "Layer fuori elenco studio (prontuario 7)", f"{lay}: {n} oggetti{nota}")

    # V02 / V03 - colore e spessore forzati sugli oggetti dello spazio modello
    col = defaultdict(int)
    lw = defaultdict(int)
    for e in msp:
        if not e.dxf.hasattr("layer"):
            continue
        c = e.dxf.get("color", 256)
        if c not in (256, 0) or e.dxf.hasattr("true_color"):
            col[e.dxf.layer] += 1
        w = e.dxf.get("lineweight", -1)
        if w not in (-1, -2):
            lw[e.dxf.layer] += 1
    for lay, n in sorted(col.items()):
        add("V02", "ERRORE", "Colore forzato: deve essere DaLayer, decide lo stile di stampa (prontuario 8)",
            f"{lay}: {n} oggetti")
    for lay, n in sorted(lw.items()):
        add("V03", "AVVISO", "Spessore forzato: deve essere DaLayer (prontuario 8)", f"{lay}: {n} oggetti")

    # V04 / V05 - testi
    fattore = MM_PER_UNITA[unita] / float(scala)
    for e in msp.query("TEXT MTEXT"):
        rot = e.dxf.get("rotation", 0.0) % 360
        testo = (e.dxf.text if e.dxftype() == "TEXT" else e.plain_text())[:40].replace("\n", " ")
        if 90.5 < rot < 359.5:
            add("V04", "ERRORE", "Testo non leggibile dal basso o da destra (CEI EN 61082-1, 5.2)",
                f"'{testo}' ruotato {rot:.0f} gradi, layer {e.dxf.layer}")
        elif 0.5 < rot < 89.5:
            add("V04", "AVVISO", "Testo obliquo (CEI EN 61082-1, 5.2)", f"'{testo}' ruotato {rot:.0f} gradi")
        h = e.dxf.height if e.dxftype() == "TEXT" else e.dxf.get("char_height", 0)
        if stile_annotativo(doc, e.dxf.get("style", "Standard")):
            continue
        carta = h * fattore
        if 0 < carta < ALTEZZA_MIN_MM - 0.01:
            add("V05", "AVVISO", f"Testo sotto {ALTEZZA_MIN_MM} mm su carta (prontuario 7)",
                f"'{testo}' = {carta:.2f} mm (h {h:g} in {unita}, 1:{scala:g}), layer {e.dxf.layer}")

    # V06 - stili di stampa dei layout
    for lay in doc.layouts:
        if lay.name == "Model":
            continue
        ctb = (lay.dxf_layout.dxf.get("current_style_sheet") or "").upper()
        if len(lay) and ctb not in CTB_AMMESSI:
            add("V06", "AVVISO", "Stile di stampa non PEMAJ (prontuario 7)", f"layout {lay.name}: {ctb or 'nessuno'}")

    # V07 / V08 - cartiglio
    cartiglio = None
    for lay in doc.layouts:
        for i in lay.query("INSERT"):
            n = i.dxf.name.upper()
            if n == "PEMAJ_CARTIGLIO_BASE" or "CARTIGLIO" in n:
                cartiglio = (lay, i)
                break
        if cartiglio:
            break
    if not cartiglio:
        add("V07", "ERRORE", "Frontalino assente (prontuario 7)", "nessun CARTIGLIO BASE / PEMAJ_CARTIGLIO_BASE")
    else:
        lay, ins = cartiglio
        att = {a.dxf.tag: a.dxf.text for a in ins.attribs}
        if att:
            for tag, val in att.items():
                if val.strip() in SEGNAPOSTO:
                    add("V07", "ERRORE", "Campo del frontalino non compilato", f"{tag} = '{val}'")
            sc = att.get("SCALA", "")
            vps = [v for v in lay.query("VIEWPORT") if v.dxf.id > 1 and v.dxf.height]
            if sc.startswith("1:") and vps:
                fatt_vp = vps[0].dxf.view_height / vps[0].dxf.height     # unita' modello per mm di carta
                reale = fatt_vp * MM_PER_UNITA[unita]
                try:
                    dichiarata = float(sc[2:].replace(",", "."))
                    if abs(reale - dichiarata) / dichiarata > 0.02:
                        add("V08", "AVVISO", "Scala del cartiglio diversa dalla finestra",
                            f"cartiglio {sc}, finestra 1:{reale:g} (unita' {unita})")
                except ValueError:
                    pass
        else:
            add("V07", "INFO", "Frontalino come riferimento esterno: campi non leggibili dal DXF",
                f"{ins.dxf.name} nel layout {lay.name}")

    # V09 - legenda DI.RI.
    if tipo == "diri":
        tutti = " ".join((t.dxf.text if t.dxftype() == "TEXT" else t.plain_text()).upper()
                         for l in doc.layouts for t in l.query("TEXT MTEXT"))
        if "ESISTENTE" not in tutti:
            add("V09", "ERRORE", "Tavola DI.RI.: manca in legenda la convenzione nero = esistente, colore = nuovo",
                "nessun testo 'ESISTENTE'")
    return trovati


def scrivi_report(trovati, dest, tavola):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Verifica"
    ws.append(["Codice", "Gravita'", "Regola", "Dettaglio"])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    colori = {"ERRORE": "C00000", "AVVISO": "B36B00", "INFO": "595959"}
    for t in trovati[:5000]:
        ws.append(list(t))
        for c in ws[ws.max_row]:
            c.font = Font(name="Arial", size=10, color=colori.get(t[1], "000000"))
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", [8, 10, 60, 80]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Riepilogo")
    ws2.append(["Tavola", str(tavola)])
    for (cod, grav), n in sorted(Counter((t[0], t[1]) for t in trovati).items()):
        ws2.append([cod, grav, n])
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tavola")
    ap.add_argument("--unita", default="cm", choices=list(MM_PER_UNITA))
    ap.add_argument("--scala", type=float, default=100.0)
    ap.add_argument("--tipo", default="progetto", choices=["progetto", "diri", "schema"])
    ap.add_argument("--progetto")
    a = ap.parse_args()
    p = Path(a.tavola)
    if not p.exists():
        esci({"status": "errore", "code": "TAVOLA_NON_TROVATA", "fix": f"Controllare il percorso: {p}"}, False)
    with open(p, "rb") as f:
        if f.read(6).startswith(b"AC10"):
            esci({"status": "errore", "code": "E_UN_DWG",
                  "fix": "Il file e' un DWG: in LT SALVACOME -> AutoCAD 2018 DXF e rilanciare."}, False)
    try:
        trovati = verifica(p, a.unita, a.scala, a.tipo)
    except Exception as ex:
        esci({"status": "errore", "code": "DXF_ILLEGGIBILE", "message": str(ex),
              "fix": "Aprire in LT, comando RECUPERA, poi SALVACOME DXF 2018."}, False)
    if a.progetto:
        dest = BASE.parent.parent / "Projects" / a.progetto / f"{p.stem}_verifica.xlsx"
    else:
        dest = OUT / "verifiche" / f"{p.stem}_verifica.xlsx"
    scrivi_report(trovati, dest, p)
    conta = Counter(t[1] for t in trovati)
    esci({"status": "ok" if not conta.get("ERRORE") else "errori",
          "tavola": str(p), "report": str(dest),
          "errori": conta.get("ERRORE", 0), "avvisi": conta.get("AVVISO", 0),
          "per_codice": dict(Counter(t[0] for t in trovati)),
          "esempi": [list(t) for t in trovati[:15]]}, True)


if __name__ == "__main__":
    main()
