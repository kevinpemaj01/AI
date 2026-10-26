"""
Studio Pemaj - Prepara la base architettonica: layer e colori dello studio, file snello.

Uso:
  python prepara_base.py "<pianta.dxf>" [--progetto <P>] [--viste 1,3] [--mappa <tabella.xlsx>]
                         [--unita-ingresso auto|mm|cm|m] [--unita-uscita mm|cm] [--solo-analisi]

Cosa fa (in ordine):
  1. legge il DXF dell'architetto (un DWG va prima salvato come DXF; se c'e' dwg2dxf lo converte da se')
  2. ogni layer va su un layer dello studio: MURI, INF, ARREDI, SANIT, MACCHINE, SCRITTE oppure ELIMINA
     (regole sul nome del layer + mappa memorizzata dello studio riferimenti/mappa_layer.json)
  3. elimina: quote, retini, campiture, perimetri delle aree, immagini, riferimenti esterni, oggetti 3D,
     punti, direttrici, cartigli/cornici dell'architetto, testi che non sono nomi di locali
     (misure delle finestre, numeri, superfici, note lunghe), testi dentro i blocchi
  4. tiene: muri, porte, finestre, scale, arredi, macchinari, nomi dei locali (layer SCRITTE)
  5. colore, spessore e trasparenza di ogni oggetto = DaLayer (decide il CTB PEMAJ)
  6. riconosce le viste del foglio (piante, sezioni, prospetti, cartiglio): di default tiene le PIANTE
     (o tutto il disegno se non ci sono titoli); --viste per scegliere
  7. unita' in mm (o cm), quote Z a zero, oggetti doppi e linee nulle eliminati
  8. scrive un DXF nuovo con solo cio' che serve (blocchi, layer, stili non usati non vengono copiati)
     + tabella dei layer (Excel) da confermare + anteprime prima/dopo (PNG) + riepilogo JSON

Output: Projects/<P>/<nome>_base.dxf (o out/basi/), <nome>_base_layer.xlsx, <nome>_base_viste.png,
        <nome>_base_dopo.png. Il DXF si apre in LT e si salva come DWG.

La tabella: foglio Layer e foglio Blocchi, colonna "Destinazione" modificabile (MURI, INF, INFISSI, ARREDI,
SANIT, MACCHINE, SCRITTE, ELIMINA; solo per i layer anche MISTO = linee dritte >= 3 m su MURI, il resto su INF).
Blocchi: sul layer dei muri o complessi (macchinari) -> MACCHINE. Rilanciando con --mappa <tabella.xlsx> le scelte vengono applicate e memorizzate
in riferimenti/mappa_layer.json, cosi' la volta dopo lo stesso layer va da solo al posto giusto.
"""
import argparse
import json
import math
import os
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

import ezdxf
from ezdxf import recover
from ezdxf.math import Matrix44

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comune import BASE, RIF, OUT, esci  # noqa: E402

MAPPA_STUDIO = RIF / "mappa_layer.json"
DESTINAZIONI = ("MURI", "INF", "INFISSI", "ARREDI", "SANIT", "MACCHINE", "SCRITTE", "ELIMINA", "MISTO")
# MISTO = layer che mescola muri e altro: linee e polilinee dritte lunghe almeno MISTO_MURO_MM -> MURI, il resto -> INF
MISTO_MURO_MM = 3000
DEFAULT = "INF"            # layer sconosciuti: grigio sottile, come la pianta di base nelle tavole E01

# regole sul nome del layer: la prima che corrisponde vince. Parole ASCII = inizio di una parola del nome;
# parole non ASCII (layer di programmi cinesi) = contenute nel nome.
REGOLE = [
    ("ELIMINA", ["quot", "dim", "cot", "misur", "superf", "perimetr", "aree", "area", "retin", "hatch",
                 "campit", "riempiment", "center", "centre", "assi", "asse", "axis", "cartigl", "testalin",
                 "cornic", "squadrat", "titol", "defpoints", "viewport", "vport", "sezion", "section",
                 "尺寸", "剖面", "中心", "标注"]),
    ("MURI", ["mur", "wall", "paret", "tramez", "pilastr", "struttur", "tampon"]),
    ("INF", ["port", "finestr", "serrament", "infiss", "door", "window", "vetr", "glass", "soglie",
             "davanzal", "scal", "stair", "inf"]),
    ("SANIT", ["sanit", "wc", "bagn", "lavab"]),
    ("ARREDI", ["arred", "furn", "mobil"]),
    ("MACCHINE", ["macchin", "machine", "manufatt", "attrezz", "equip", "impiant"]),
    ("SCRITTE", ["testi", "text", "scritt", "nomi"]),
]

TIPI_ELIMINATI = {
    "DIMENSION", "ARC_DIMENSION", "LARGE_RADIAL_DIMENSION", "LEADER", "MULTILEADER", "MLEADER",
    "HATCH", "MPOLYGON", "SOLID", "TRACE", "WIPEOUT", "IMAGE", "OLE2FRAME", "OLEFRAME", "POINT",
    "REGION", "3DSOLID", "BODY", "SURFACE", "MESH", "ACAD_TABLE", "ACAD_PROXY_ENTITY", "TOLERANCE",
    "RAY", "XLINE", "PDFUNDERLAY", "DWFUNDERLAY", "DGNUNDERLAY", "UNDERLAY", "HELIX", "LIGHT",
    "SECTION", "VIEWPORT", "ATTDEF", "3DFACE", "POLYFACE", "PLANESURFACE", "EXTRUDEDSURFACE",
    "LOFTEDSURFACE", "REVOLVEDSURFACE", "SWEPTSURFACE", "NURBSURFACE", "SHAPE",
}
UNITA_MM = {"mm": 1.0, "cm": 10.0, "m": 1000.0}
INSUNITS_NOME = {4: "mm", 5: "cm", 6: "m"}
TITOLI = re.compile(r"pianta|piano|planimetr|layout|prospett|sezion|stato di fatto|copertur", re.I)
TITOLI_PIANTA = re.compile(r"pianta|planimetr|layout", re.I)
NON_LOCALE = re.compile(r"(\bmq\b|m2|m²|\bsup\b|superf|\bs\s*[=:]|\bh\s*[=:]|\bquota\b|\bscala\b|"
                        r"\btav\b|tavola|\bdata\b|committ|progett|\barch\b|\bing\b|\bgeom\b|\bdott\b)", re.I)
GRIGLIA = 300            # celle sul lato lungo per riconoscere le viste
DILATA = 4               # celle di tolleranza tra parti della stessa vista (~1,3% del foglio)


# ---------------------------------------------------------------- lettura
def leggi(path):
    path = Path(path)
    with open(path, "rb") as f:
        dwg = f.read(6).startswith(b"AC10")
    if dwg:
        conv = os.environ.get("PEMAJ_DWG2DXF") or shutil.which("dwg2dxf")
        if not conv:
            esci({"status": "errore", "code": "E_UN_DWG",
                  "fix": "Il file e' un DWG: in LT aprirlo e SALVACOME -> AutoCAD 2018 DXF, poi rilanciare."}, False)
        dxf = path.with_suffix(".convertito.dxf")
        subprocess.run([conv, "-y", "-o", str(dxf), str(path)], capture_output=True, timeout=1800)
        path = dxf
    doc, auditor = recover.readfile(str(path))
    return doc


# ---------------------------------------------------------------- regole
def _parole(nome):
    return [p for p in re.split(r"[^0-9a-zà-ù]+", nome.lower()) if p]


def destinazione_layer(nome, memoria):
    """(destinazione, motivo) per un layer dell'architetto."""
    if nome.lower() in memoria:
        return memoria[nome.lower()], "mappa dello studio"
    parole = _parole(nome)
    for dest, chiavi in REGOLE:
        for k in chiavi:
            if k.isascii():
                if any(p.startswith(k) for p in parole):
                    return dest, f"nome contiene '{k}'"
            elif k in nome:
                return dest, f"nome contiene '{k}'"
    if nome == "0":
        return DEFAULT, "layer 0"
    return DEFAULT, "non riconosciuto: grigio come base"


def testo_di(e):
    try:
        return e.dxf.text if e.dxftype() in ("TEXT", "ATTRIB") else e.plain_text()
    except Exception:
        return ""


def e_nome_locale(s):
    s = " ".join(s.split())
    lettere = sum(c.isalpha() for c in s)
    return 3 <= lettere and len(s) <= 40 and not NON_LOCALE.search(s)


def carica_memoria():
    try:
        return {k.lower(): v for k, v in json.loads(MAPPA_STUDIO.read_text(encoding="utf-8")).items()}
    except (OSError, ValueError):
        return {}


def leggi_tabella(xlsx):
    """Scelte dell'utente: foglio Layer (nome layer) e foglio Blocchi (chiave 'blocco:<nome>')."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, data_only=True)
    out = {}
    for foglio, col, pref in (("Layer", "layer", ""), ("Blocchi", "blocco", "blocco:")):
        if foglio not in wb.sheetnames:
            continue
        righe = list(wb[foglio].iter_rows(values_only=True))
        h = [str(c or "").strip().lower() for c in righe[0]]
        il, idd = h.index(col), h.index("destinazione")
        for r in righe[1:]:
            if r[il] is None:
                continue
            d = str(r[idd] or "").strip().upper()
            if d not in DESTINAZIONI:
                esci({"status": "errore", "code": "DESTINAZIONE_NON_VALIDA",
                      "message": f"{col} {r[il]}: '{d}'", "fix": f"Usare uno di {', '.join(DESTINAZIONI)}"}, False)
            out[pref + str(r[il]).lower()] = d
    return out


def e_muro_lungo(e, f):
    """Per i layer MISTO: tratto dritto lungo (in mm dopo la conversione di unita') = muro."""
    t = e.dxftype()
    try:
        if t == "LINE":
            return (e.dxf.end - e.dxf.start).magnitude * f >= MISTO_MURO_MM
        if t == "LWPOLYLINE":
            p = e.get_points("xyb")
            if any(abs(q[2]) > 1e-6 for q in p):
                return False
            return any(math.hypot(b[0] - a[0], b[1] - a[1]) * f >= MISTO_MURO_MM for a, b in zip(p, p[1:]))
    except Exception:
        return False
    return False


# ---------------------------------------------------------------- blocchi
CURVE = {"ARC", "CIRCLE", "SPLINE", "ELLIPSE"}
SOGLIA_MACCHINA = 150        # oggetti in un blocco oltre i quali e' un macchinario, non una porta o un arredo
SOGLIA_CURVE = (40, 0.3)     # oppure almeno 40 oggetti con il 30% di curve (disegni meccanici)


def statistiche_blocco(doc, nome, cache, prof=0):
    if nome in cache:
        return cache[nome]
    n = curve = 0
    b = doc.blocks.get(nome)
    if b is not None and prof < 8:
        for e in b:
            t = e.dxftype()
            if t == "INSERT":
                n2, c2 = statistiche_blocco(doc, e.dxf.name, cache, prof + 1)
                n, curve = n + n2, curve + c2
            else:
                n += 1
                curve += t in CURVE
    cache[nome] = (n, curve)
    return cache[nome]


def destinazione_blocco(doc, nome, dest_layer, memoria, cache):
    """Un blocco complesso (macchinario) va su MACCHINE anche se e' inserito sul layer dei muri o sul layer 0.
    Porte, finestre, arredi e sanitari seguono il layer. La scelta dell'utente (foglio Blocchi) vince."""
    k = "blocco:" + nome.lower()
    if k in memoria:
        return memoria[k], "mappa dello studio"
    n, c = statistiche_blocco(doc, nome, cache)
    if dest_layer in ("MURI", "MISTO") and not re.search(r"mur|wall|pilastr|colonn|pillar", nome, re.I):
        return "MACCHINE", f"blocco sul layer dei muri ({n} oggetti): i muri non sono blocchi"
    if dest_layer in ("INF", "INFISSI", "SCRITTE") and (
            n >= SOGLIA_MACCHINA or (n >= SOGLIA_CURVE[0] and c >= SOGLIA_CURVE[1] * n)):
        return "MACCHINE", f"blocco complesso ({n} oggetti, {c} curve)"
    if dest_layer == "MISTO":
        return "INF", "blocco su layer misto"
    return dest_layer, "come il layer"


# ---------------------------------------------------------------- unita'
def stima_unita(doc, punti, testi_h):
    """Unita' del disegno: la pianta deve misurare 3 m - 1,5 km e i testi 5 cm - 2 m."""
    import numpy as np
    if not punti:
        return "mm", "nessuna geometria"
    p = np.array(punti)
    lato = max(np.percentile(p[:, 0], 99) - np.percentile(p[:, 0], 1),
               np.percentile(p[:, 1], 99) - np.percentile(p[:, 1], 1))
    h = float(np.median(testi_h)) if testi_h else None
    voti = {}
    for u, f in UNITA_MM.items():
        ok_lato = 3000 <= lato * f <= 1_500_000
        ok_testo = h is None or 50 <= h * f <= 2000
        voti[u] = ok_lato * 2 + ok_testo
    dich = INSUNITS_NOME.get(doc.header.get("$INSUNITS", 0))
    migliore = max(voti.values())
    cand = [u for u, v in voti.items() if v == migliore]
    u = dich if dich in cand else ("mm" if "mm" in cand else cand[0])
    nota = f"lato disegno {lato:.4g} unita', testo medio {h:.3g}" if h else f"lato disegno {lato:.4g} unita'"
    if dich and dich != u:
        nota += f"; il file dichiara {dich} ma le misure dicono {u}"
    return u, nota


# ---------------------------------------------------------------- viste
def punti_entita(e):
    t = e.dxftype()
    try:
        if t == "LINE":
            return [e.dxf.start, e.dxf.end]
        if t == "LWPOLYLINE":
            pts = e.get_points("xy")
            return pts[:: max(1, len(pts) // 20)] + pts[-1:]
        if t in ("CIRCLE", "ARC", "ELLIPSE"):
            return [e.dxf.center]
        if t == "INSERT":           # ingombro reale: il punto base di un blocco puo' essere lontano dal disegno
            from ezdxf import bbox
            b = bbox.extents([e], fast=True)
            if b.has_data:
                return [b.extmin, b.extmax, b.center]
            return [e.dxf.insert]
        if t in ("TEXT", "MTEXT"):
            return [e.dxf.insert]
        if t == "SPLINE":
            cp = e.control_points if len(e.control_points) else e.fit_points
            return [cp[0], cp[-1]] if len(cp) else []
        if t == "POLYLINE":
            v = [x.dxf.location for x in e.vertices]
            return v[:: max(1, len(v) // 20)]
    except Exception:
        return []
    return []


def trova_viste(ents):
    """Raggruppa le entita' in viste (piante, sezioni, prospetti, cartiglio) per vicinanza."""
    import numpy as np
    pe = []
    testo = []
    for i, e in enumerate(ents):
        for p in punti_entita(e):
            pe.append((p[0], p[1], i))
            testo.append(e.dxftype() in ("TEXT", "MTEXT"))
    if not pe:
        return [], {}
    a = np.array(pe)
    x0, x1 = np.percentile(a[:, 0], [0.5, 99.5])
    y0, y1 = np.percentile(a[:, 1], [0.5, 99.5])
    mx, my = (x1 - x0) * 0.1 + 1e-9, (y1 - y0) * 0.1 + 1e-9
    x0, x1, y0, y1 = x0 - mx, x1 + mx, y0 - my, y1 + my
    cel = max(x1 - x0, y1 - y0) / GRIGLIA
    nx, ny = int((x1 - x0) / cel) + 1, int((y1 - y0) / cel) + 1
    dentro = (a[:, 0] >= x0) & (a[:, 0] <= x1) & (a[:, 1] >= y0) & (a[:, 1] <= y1)
    ix = ((a[:, 0] - x0) / cel).astype(int).clip(0, nx - 1)
    iy = ((a[:, 1] - y0) / cel).astype(int).clip(0, ny - 1)
    occ = np.zeros((nx, ny), bool)
    geo = dentro & ~np.array(testo)        # i testi (titoli, note) non devono unire due viste vicine
    occ[ix[geo], iy[geo]] = True
    dil = occ.copy()
    for dx in range(-DILATA, DILATA + 1):
        for dy in range(-DILATA, DILATA + 1):
            dil |= np.roll(np.roll(occ, dx, 0), dy, 1)
    lab = np.full((nx, ny), -1, int)
    n = 0
    for sx, sy in zip(*np.nonzero(dil)):
        if lab[sx, sy] >= 0:
            continue
        pila = [(sx, sy)]
        lab[sx, sy] = n
        while pila:
            cx, cy = pila.pop()
            for vx, vy in ((cx + 1, cy), (cx - 1, cy), (cx, cy + 1), (cx, cy - 1)):
                if 0 <= vx < nx and 0 <= vy < ny and dil[vx, vy] and lab[vx, vy] < 0:
                    lab[vx, vy] = n
                    pila.append((vx, vy))
        n += 1
    vista_di = {}
    voti = defaultdict(Counter)
    for k in range(len(a)):
        i = int(a[k, 2])
        voti[i][int(lab[ix[k], iy[k]]) if dentro[k] else -1] += 1
    for i, c in voti.items():
        vista_di[i] = c.most_common(1)[0][0]
        if -1 in c and len(c) > 1 and c[-1] >= sum(c.values()) / 2:
            vista_di[i] = -1           # meta' o piu' dell'oggetto fuori dal disegno: e' fuori
    viste = []
    for v in range(n):
        sel = (lab[ix, iy] == v) & dentro
        if not sel.any():
            continue
        viste.append({"id": v, "punti": int(sel.sum()),
                      "bbox": [float(a[sel, 0].min()), float(a[sel, 1].min()),
                               float(a[sel, 0].max()), float(a[sel, 1].max())]})
    return viste, vista_di


def nomina_e_scegli(viste, ents, vista_di, scelte):
    """Titoli (PIANTA PIANO TERRA, PROSPETTO...) e viste da tenere."""
    titoli = []
    hs = [e.dxf.char_height if e.dxftype() == "MTEXT" else e.dxf.height
          for e in ents if e.dxftype() in ("TEXT", "MTEXT")]
    hmed = sorted(hs)[len(hs) // 2] if hs else 0
    for i, e in enumerate(ents):
        if e.dxftype() in ("TEXT", "MTEXT"):
            s = " ".join(testo_di(e).split())
            h = e.dxf.char_height if e.dxftype() == "MTEXT" else e.dxf.height
            if TITOLI.search(s) and len(s) <= 50 and h >= hmed:
                titoli.append((s, e.dxf.insert.x, e.dxf.insert.y))
    tot = sum(v["punti"] for v in viste) or 1
    diag = lambda b: math.hypot(b[2] - b[0], b[3] - b[1])  # noqa: E731
    dmax = max((diag(v["bbox"]) for v in viste), default=1)
    for v in viste:
        v["titolo"] = ""
    for s, x, y in titoli:          # ogni titolo va alla vista grande piu' vicina
        best, dbest = None, None
        for v in viste:
            if v["punti"] < 0.01 * tot:
                continue
            b = v["bbox"]
            d = math.hypot(max(b[0] - x, 0, x - b[2]), max(b[1] - y, 0, y - b[3]))
            if d > 0.15 * diag(b) + 1e-9:
                continue
            if y < b[1]:            # titolo sotto la vista: vale meno di un titolo sopra
                d *= 1.5
            if dbest is None or d < dbest:
                best, dbest = v, d
        if best is not None and not best["titolo"]:
            best["titolo"] = s
    ci_sono_piante = any(TITOLI_PIANTA.search(v["titolo"]) for v in viste)
    for v in viste:
        if scelte is not None:
            v["tieni"] = v["id"] in scelte
        elif ci_sono_piante:
            v["tieni"] = bool(TITOLI_PIANTA.search(v["titolo"]))
        else:
            v["tieni"] = v["punti"] >= 0.01 * tot or diag(v["bbox"]) >= 0.05 * dmax
    # parti senza titolo dentro una vista tenuta, o appena sotto/sopra e allineate: stessa vista
    tenute = [v for v in viste if v["tieni"]]
    for v in viste:
        if v["tieni"] or scelte is not None or v["titolo"]:
            continue
        b = v["bbox"]
        for t in tenute:
            T = t["bbox"]
            m = 0.05 * diag(T)
            dentro = (b[0] >= T[0] - m and b[2] <= T[2] + m and b[1] >= T[1] - m and b[3] <= T[3] + m
                      and v["punti"] < 0.2 * t["punti"])
            sovrap = min(b[2], T[2]) - max(b[0], T[0])
            gap = max(T[1] - b[3], b[1] - T[3])
            attaccata = (sovrap >= 0.7 * (b[2] - b[0]) and 0 <= gap < 0.15 * (T[3] - T[1]))
            if dentro or attaccata:
                v["tieni"] = True
                v["titolo"] = f"(parte di {t['id']})"
                break
    return viste


# ---------------------------------------------------------------- pulizia
def normalizza(e):
    e.dxf.color = 256
    for k in ("true_color", "transparency", "color_name"):
        if e.dxf.hasattr(k):
            e.dxf.discard(k)
    e.dxf.lineweight = -1
    lt = (e.dxf.get("linetype") or "BYLAYER").upper()
    if lt in ("CONTINUOUS", "BYBLOCK"):
        e.dxf.linetype = "BYLAYER"


def pulisci_blocchi(doc, mappa):
    """Dentro le definizioni dei blocchi: stessi layer dello studio, via quote/retini/testi."""
    tolti = Counter()
    for blk in doc.blocks:
        if blk.name.lower().startswith(("*model_space", "*paper_space")) or blk.block_record.is_any_layout:
            continue
        if blk.block.is_xref:
            continue
        for e in list(blk):
            t = e.dxftype()
            lay = e.dxf.get("layer", "0")
            dest = None if lay == "0" else mappa.get(lay, (DEFAULT, ""))[0]
            if t in TIPI_ELIMINATI or t in ("TEXT", "MTEXT") or dest == "ELIMINA":
                blk.delete_entity(e)
                tolti[t] += 1
                continue
            if dest:
                e.dxf.layer = "INF" if dest == "MISTO" else dest
            normalizza(e)
    return tolti


def chiave_doppio(e, r=0.5):
    q = lambda v: round(v / r)  # noqa: E731
    t = e.dxftype()
    if t == "LINE":
        a, b = (q(e.dxf.start.x), q(e.dxf.start.y)), (q(e.dxf.end.x), q(e.dxf.end.y))
        return (t, e.dxf.layer) + tuple(sorted((a, b)))
    if t == "LWPOLYLINE":
        return (t, e.dxf.layer, e.closed) + tuple((q(p[0]), q(p[1])) for p in e.get_points("xy"))
    if t in ("CIRCLE", "ARC"):
        c = e.dxf.center
        ang = (round(e.dxf.start_angle, 1), round(e.dxf.end_angle, 1)) if t == "ARC" else (0, 0)
        return (t, e.dxf.layer, q(c.x), q(c.y), q(e.dxf.radius)) + ang
    if t == "INSERT":
        p = e.dxf.insert
        return (t, e.dxf.layer, e.dxf.name, q(p.x), q(p.y), round(e.dxf.rotation, 2),
                round(e.dxf.xscale, 4), round(e.dxf.yscale, 4))
    return None


def appiattisci(e):
    t = e.dxftype()
    try:
        if t == "LINE":
            s, f = e.dxf.start, e.dxf.end
            e.dxf.start, e.dxf.end = (s.x, s.y, 0), (f.x, f.y, 0)
        elif t == "LWPOLYLINE":
            e.dxf.elevation = 0
        elif t in ("CIRCLE", "ARC", "ELLIPSE"):
            c = e.dxf.center
            e.dxf.center = (c.x, c.y, 0)
        elif t in ("INSERT", "TEXT", "MTEXT"):
            p = e.dxf.insert
            e.dxf.insert = (p.x, p.y, 0)
    except Exception:
        pass


# ---------------------------------------------------------------- anteprime
def anteprima(ents_o_doc, dest, viste=None, titolo=""):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.collections import LineCollection
        from matplotlib.patches import Rectangle
    except ImportError:
        return None
    from array import array
    import numpy as np
    from ezdxf import path as dxfpath
    seg = defaultdict(lambda: array("f"))       # x1 y1 x2 y2 in float32: poca memoria anche con milioni di linee

    def aggiungi(lay, pts):
        a = seg[lay]
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            if abs(x2 - x1) + abs(y2 - y1) >= minimo:     # tratti invisibili all'anteprima: saltati
                a.extend((x1 - ox, y1 - oy, x2 - ox, y2 - oy))

    def disegna(v, lay):
        t = v.dxftype()
        if t == "LINE":
            aggiungi(lay, [(v.dxf.start.x, v.dxf.start.y), (v.dxf.end.x, v.dxf.end.y)])
        elif t in ("LWPOLYLINE", "POLYLINE", "ARC", "CIRCLE", "ELLIPSE", "SPLINE"):
            aggiungi(lay, [(q.x, q.y) for q in dxfpath.make_path(v).flattening(distance=passo)])   # WCS

    ents = list(ents_o_doc)
    from ezdxf import bbox as _bb
    ext = _bb.extents([e for e in ents if e.dxftype() == "LINE"][:20000] or ents[:2000], fast=True)
    passo = max(ext.size.x, ext.size.y) / 2000 if ext.has_data else 1.0
    minimo = passo / 2
    ox, oy = (ext.extmin.x, ext.extmin.y) if ext.has_data else (0.0, 0.0)   # origine locale per float32
    for e in ents:
        try:
            t = e.dxftype()
            if t == "INSERT":
                for v in e.virtual_entities():
                    if v.dxftype() != "INSERT":
                        disegna(v, e.dxf.layer)
            elif t in ("TEXT", "MTEXT"):
                p = e.dxf.insert
                h = e.dxf.char_height if t == "MTEXT" else e.dxf.height
                seg["__TESTI"].extend((p.x - ox, p.y - oy, p.x + h * 4 - ox, p.y - oy))
            else:
                disegna(e, e.dxf.layer)
        except Exception:
            continue
    colori = {"MURI": "#3355ff", "INF": "#999999", "INFISSI": "#999999", "ARREDI": "#e08a00",
              "SANIT": "#b39b00", "MACCHINE": "#cc33cc", "SCRITTE": "#000000", "__TESTI": "#000000"}
    fig = plt.figure(figsize=(16, 11), dpi=90)
    ax = fig.add_axes([0.01, 0.01, 0.98, 0.95])
    for lay, a in seg.items():
        if len(a):
            ax.add_collection(LineCollection(np.frombuffer(a, dtype=np.float32).reshape(-1, 2, 2),
                                             linewidths=0.4, colors=colori.get(lay, "#999999")))
    seg.clear()
    if viste:
        for v in viste:
            b = [v["bbox"][0] - ox, v["bbox"][1] - oy, v["bbox"][2] - ox, v["bbox"][3] - oy]
            ax.add_patch(Rectangle((b[0], b[1]), b[2] - b[0], b[3] - b[1], fill=False,
                                   ec="#00aa00" if v["tieni"] else "#dd0000", lw=1.5, ls="-" if v["tieni"] else "--"))
            ax.text(b[0], b[3], f" {v['id']} {v['titolo'][:30]}", color="#00aa00" if v["tieni"] else "#dd0000",
                    fontsize=9, va="bottom")
    ax.autoscale()
    ax.set_aspect("equal")
    ax.set_title(titolo)
    ax.axis("off")
    fig.savefig(dest)
    plt.close(fig)
    return str(dest)


# ---------------------------------------------------------------- tabella layer
def scrivi_tabella(righe, dest, righe_blocchi=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Layer"
    ws.append(["Layer", "Oggetti", "Tipi", "Proposta", "Motivo", "Destinazione"])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for r in righe:
        ws.append([r["layer"], r["oggetti"], r["tipi"], r["proposta"], r["motivo"], r["proposta"]])
        ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor="DDEBF7")
    dv = DataValidation(type="list", formula1='"' + ",".join(DESTINAZIONI) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{max(2, ws.max_row)}")
    for col, w in zip("ABCDEF", [34, 10, 50, 12, 34, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb2 = wb.create_sheet("Blocchi")
    wb2.append(["Blocco", "Inserimenti", "Oggetti nel blocco", "Layer di inserimento", "Proposta", "Motivo",
                "Destinazione"])
    for c in wb2[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for r in righe_blocchi:
        wb2.append([r["blocco"], r["inserimenti"], r["oggetti"], r["layer"], r["proposta"], r["motivo"],
                    r["proposta"]])
        wb2.cell(wb2.max_row, 7).fill = PatternFill("solid", fgColor="DDEBF7")
    dv2 = DataValidation(type="list", formula1='"' + ",".join(DESTINAZIONI) + '"', allow_blank=False)
    wb2.add_data_validation(dv2)
    dv2.add(f"G2:G{max(2, wb2.max_row)}")
    for col, w in zip("ABCDEFG", [34, 12, 16, 30, 12, 34, 14]):
        wb2.column_dimensions[col].width = w
    wb2.freeze_panes = "A2"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


# ---------------------------------------------------------------- principale
def prepara(src, out_dir, viste_scelte=None, tabella=None, unita_in="auto", unita_out="mm",
            solo_analisi=False):
    import genera_template
    memoria = carica_memoria()
    if tabella:
        memoria.update(leggi_tabella(tabella))
    doc = leggi(src)
    msp = doc.modelspace()
    nome = Path(src).stem
    rep = {"file": str(src), "oggetti_prima": len(msp)}

    # 1. layer -> destinazione
    uso = defaultdict(Counter)
    for e in msp:
        uso[e.dxf.get("layer", "0")][e.dxftype()] += 1
    mappa = {lay.dxf.name: destinazione_layer(lay.dxf.name, memoria) for lay in doc.layers}
    for lay in uso:
        mappa.setdefault(lay, destinazione_layer(lay, memoria))
    righe = [{"layer": l, "oggetti": sum(c.values()), "tipi": ", ".join(f"{t} {n}" for t, n in c.most_common(5)),
              "proposta": mappa[l][0], "motivo": mappa[l][1]} for l, c in sorted(uso.items(), key=lambda kv: -sum(kv[1].values()))]

    # 2. cosa resta nello spazio modello
    tolti = Counter()
    tenuti, testi_h, punti = [], [], []
    for e in msp:
        t = e.dxftype()
        dest = mappa.get(e.dxf.get("layer", "0"), (DEFAULT, ""))[0]
        if t in TIPI_ELIMINATI or dest == "ELIMINA":
            tolti[f"{t} (tipo/layer da eliminare)"] += 1
            continue
        if t == "INSERT":
            b = doc.blocks.get(e.dxf.name)
            if b is None or b.block.is_xref:
                tolti["riferimenti esterni"] += 1
                continue
        if t in ("TEXT", "MTEXT"):
            if not e_nome_locale(testo_di(e)):
                tolti["testi non nomi di locali"] += 1
                continue
            testi_h.append(e.dxf.char_height if t == "MTEXT" else e.dxf.height)
        tenuti.append(e)
    for e in tenuti:
        punti += [(p[0], p[1]) for p in punti_entita(e)[:2]]

    # 3. viste
    viste, vista_di = trova_viste(tenuti)
    viste = nomina_e_scegli(viste, tenuti, vista_di, viste_scelte)
    tieni_v = {v["id"] for v in viste if v["tieni"]}
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rep["anteprima_viste"] = anteprima(tenuti, out_dir / f"{nome}_base_viste.png", viste,
                                       "Viste trovate (verde = tenute, rosso = scartate)")
    finali = [e for i, e in enumerate(tenuti) if i not in vista_di or vista_di[i] in tieni_v]
    tolti["fuori dalle viste tenute"] = len(tenuti) - len(finali)

    # 4. unita'
    if unita_in == "auto":
        u, nota = stima_unita(doc, punti, testi_h)
    else:
        u, nota = unita_in, "indicata dall'utente"
    f = UNITA_MM[u] / UNITA_MM[unita_out]
    rep["unita"] = {"ingresso": u, "uscita": unita_out, "fattore": f, "nota": nota}
    tot = sum(v["punti"] for v in viste) or 1
    rep["viste"] = [{"id": v["id"], "titolo": v["titolo"], "tieni": v["tieni"], "punti": v["punti"],
                     "bbox": [round(x, 2) for x in v["bbox"]]}
                    for v in viste if v["punti"] >= 0.01 * tot and not v["titolo"].startswith("(parte")]
    rep["viste_minori_scartate"] = sum(1 for v in viste if not v["tieni"] and v["punti"] < 0.01 * tot)
    cache = {}
    dest_ins = {}
    per_blocco = defaultdict(lambda: {"n": 0, "layer": Counter()})
    for e in finali:
        if e.dxftype() == "INSERT":
            dl = mappa.get(e.dxf.get("layer", "0"), (DEFAULT, ""))[0]
            dest_ins[id(e)] = destinazione_blocco(doc, e.dxf.name, dl, memoria, cache)
            pb = per_blocco[e.dxf.name]
            pb["n"] += 1
            pb["layer"][e.dxf.layer] += 1
            pb["dest"] = dest_ins[id(e)]
    righe_b = [{"blocco": b, "inserimenti": v["n"], "oggetti": statistiche_blocco(doc, b, cache)[0],
                "layer": ", ".join(f"{l} {n}" for l, n in v["layer"].most_common(3)),
                "proposta": v["dest"][0], "motivo": v["dest"][1]}
               for b, v in sorted(per_blocco.items(), key=lambda kv: -kv[1]["n"] * statistiche_blocco(doc, kv[0], cache)[0])]
    rep["blocchi_su_macchine"] = sum(1 for r in righe_b if r["proposta"] == "MACCHINE" and r["motivo"].startswith("blocco"))
    tab = out_dir / f"{nome}_base_layer.xlsx"
    scrivi_tabella(righe, tab, righe_b)
    rep["tabella_layer"] = str(tab)
    if solo_analisi:
        return rep

    # 5. pulizia di blocchi e oggetti tenuti
    tolti_blocchi = pulisci_blocchi(doc, mappa)
    for e in finali:
        t = e.dxftype()
        if t in ("TEXT", "MTEXT"):
            e.dxf.layer = "SCRITTE"
        elif t == "INSERT":
            e.dxf.layer = dest_ins[id(e)][0]
        else:
            d = mappa.get(e.dxf.get("layer", "0"), (DEFAULT, ""))[0]
            if d == "MISTO":
                d = "MURI" if e_muro_lungo(e, f) else "INF"
            e.dxf.layer = d
        if e.dxf.layer == "ELIMINA":
            continue
        normalizza(e)
        if t == "INSERT":
            for a in list(e.attribs):
                if e_nome_locale(a.dxf.text):
                    a.dxf.layer = "SCRITTE"
                    normalizza(a)
                else:
                    e.attribs.remove(a)
    n0 = len(finali)
    finali = [e for e in finali if e.dxf.layer != "ELIMINA"]
    if n0 > len(finali):
        tolti["blocchi da eliminare (foglio Blocchi)"] = n0 - len(finali)
    # 6. unita', Z, doppi
    m = Matrix44.scale(f, f, f) if f != 1 else None
    visti, unici = set(), []
    for e in finali:
        if m is not None:
            try:
                e.transform(m)
            except Exception:
                tolti["non scalabili"] += 1
                continue
        appiattisci(e)
        if e.dxftype() == "LINE" and (e.dxf.start - e.dxf.end).magnitude < 0.01:
            tolti["linee di lunghezza zero"] += 1
            continue
        k = chiave_doppio(e)
        if k is not None:
            if k in visti:
                tolti["oggetti doppi"] += 1
                continue
            visti.add(k)
        unici.append(e)

    # 7. nuovo file: solo cio' che serve, con le tabelle dello studio
    from ezdxf.addons.importer import Importer
    nuovo = ezdxf.new("R2018")
    genera_template.prepara_tabelle(nuovo)
    nuovo.header["$INSUNITS"] = 4 if unita_out == "mm" else 5
    for lay in nuovo.layers:               # i layer dello studio devono esistere anche nel sorgente
        if lay.dxf.name not in doc.layers:
            doc.layers.add(lay.dxf.name, color=lay.dxf.color, linetype="Continuous")
    imp = Importer(doc, nuovo)
    imp.import_entities(unici, nuovo.modelspace())
    imp.finalize()
    nuovo.audit()
    dxf_out = out_dir / f"{nome}_base.dxf"
    nuovo.saveas(str(dxf_out))
    import gc
    n_blocchi_prima = len([b for b in doc.blocks if not b.name.startswith("*")])
    del imp, doc, msp, finali, unici, tenuti, visti
    gc.collect()
    rep["anteprima_dopo"] = anteprima(nuovo.modelspace(), out_dir / f"{nome}_base_dopo.png", None,
                                      f"Base pulita ({unita_out}) - layer studio")
    rep.update({
        "status": "ok", "file_pulito": str(dxf_out), "oggetti_dopo": len(nuovo.modelspace()),
        "eliminati": dict(tolti), "eliminati_nei_blocchi": dict(tolti_blocchi),
        "layer_dopo": dict(Counter(e.dxf.layer for e in nuovo.modelspace())),
        "blocchi": [n_blocchi_prima,
                    len([b for b in nuovo.blocks if not b.name.startswith("*")])],
        "mb": [round(Path(src).stat().st_size / 1e6, 1), round(dxf_out.stat().st_size / 1e6, 1)],
    })
    if tabella:
        salva = carica_memoria()
        salva.update(leggi_tabella(tabella))
        MAPPA_STUDIO.write_text(json.dumps(dict(sorted(salva.items())), indent=1, ensure_ascii=False),
                                encoding="utf-8")
        rep["mappa_memorizzata"] = str(MAPPA_STUDIO)
    return rep


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pianta")
    ap.add_argument("--progetto")
    ap.add_argument("--viste", help="numeri delle viste da tenere, es. 1,3 (vedi anteprima _viste.png)")
    ap.add_argument("--mappa", help="tabella dei layer corretta (xlsx) da applicare e memorizzare")
    ap.add_argument("--unita-ingresso", default="auto", choices=["auto", "mm", "cm", "m"])
    ap.add_argument("--unita-uscita", default="mm", choices=["mm", "cm"])
    ap.add_argument("--solo-analisi", action="store_true")
    a = ap.parse_args()
    p = Path(a.pianta)
    if not p.exists():
        esci({"status": "errore", "code": "PIANTA_NON_TROVATA", "fix": f"Controllare il percorso: {p}"}, False)
    out = (BASE.parent.parent / "Projects" / a.progetto) if a.progetto else (OUT / "basi")
    viste = {int(x) for x in a.viste.split(",")} if a.viste else None
    try:
        rep = prepara(p, out, viste, a.mappa, a.unita_ingresso, a.unita_uscita, a.solo_analisi)
    except PermissionError as ex:
        esci({"status": "errore", "code": "FILE_APERTO", "message": str(ex),
              "fix": "Chiudere il file in AutoCAD LT e rilanciare"}, False)
    rep.setdefault("status", "analisi")
    esci(rep)


if __name__ == "__main__":
    main()
