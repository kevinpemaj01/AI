"""
LOCK della skill standard-pemaj. Confronta con tests/expected_results.json.
Uso: python tools/standard-pemaj/tests/run_lock.py      (esce 0 se tutti i test passano)
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

T = Path(__file__).resolve().parent
BASE = T.parent
sys.path.insert(0, str(BASE))
import ezdxf                      # noqa: E402
from ezdxf.addons import acadctb  # noqa: E402
import genera_ctb                 # noqa: E402
import genera_template            # noqa: E402
import verifica_tavola            # noqa: E402
import prepara_base               # noqa: E402
from comune import MASTER, CARTIGLIO  # noqa: E402

ATTESI = json.loads((T / "expected_results.json").read_text(encoding="utf-8"))
risultati = []


def test(nome):
    def deco(fn):
        try:
            fn()
            risultati.append((nome, True, ""))
        except AssertionError as ex:
            risultati.append((nome, False, str(ex)))
        except Exception as ex:
            risultati.append((nome, False, f"{type(ex).__name__}: {ex}"))
        return fn
    return deco


def rgb_lw(ctb, aci):
    st = ctb[aci]
    c = "OGGETTO" if st.has_object_color() else "#%02X%02X%02X" % tuple(st.color)
    return c, round(st.get_lineweight(), 3)


tmp = Path(tempfile.mkdtemp(prefix="pemaj_lock_"))


@test("TC1 CTB generati = master su 255 colori x 3 stili")
def _():
    genera_ctb.genera(MASTER, tmp)
    diff = genera_ctb.confronta(MASTER, tmp)
    assert all(not v for v in diff.values()), diff


@test("TC2 valori chiave PEMAJ_PROGETTI")
def _():
    ctb = acadctb.load(str(tmp / "PEMAJ_PROGETTI.ctb"))
    for aci, (c, lw) in ATTESI["progetti"].items():
        assert rgb_lw(ctb, int(aci)) == (c, lw), f"ACI {aci}: {rgb_lw(ctb, int(aci))} invece di {(c, lw)}"


@test("TC3 DIRI: esistente nero, parti nuove a colori")
def _():
    ctb = acadctb.load(str(tmp / "PEMAJ_DIRI.ctb"))
    for aci in ATTESI["diri_neri"]:
        assert rgb_lw(ctb, aci)[0] == "#000000", f"ACI {aci} non nero in DIRI"
    for aci, (c, lw) in ATTESI["diri_colori"].items():
        assert rgb_lw(ctb, int(aci)) == (c, lw), f"ACI {aci}: {rgb_lw(ctb, int(aci))}"


@test("TC4 frontalino e cartiglio: 24 attributi, guide non stampabili e retino grigio dietro, testi centrati")
def _():
    r = genera_template.genera(CARTIGLIO, tmp / "tpl.dxf")
    assert r["frontalino_mm"] == ATTESI["frontalino_mm"], r
    assert genera_template.controlla(tmp / "tpl.dxf") == []
    doc = ezdxf.readfile(str(tmp / "tpl.dxf"))
    tags = sorted(a.dxf.tag for a in doc.blocks.get("PEMAJ_CARTIGLIO_BASE").query("ATTDEF"))
    assert tags == sorted(ATTESI["attributi"]), tags
    blk = list(doc.blocks.get("PEMAJ_CARTIGLIO_BASE"))
    tratt = [i for i, e in enumerate(blk) if e.dxftype() == "HATCH"]
    grigi = [i for i in tratt if genera_template.e_sfondo_grigio(blk[i], doc)]
    nguide = sum(1 for e in blk if e.dxf.layer == "GUIDE")
    assert grigi and grigi == list(range(nguide, nguide + len(grigi))), \
        "il retino grigio del logo (ombra) deve stare dietro a tutto (subito dopo le guide non stampabili)"
    for a in blk:
        if a.dxftype() == "ATTDEF" and a.dxf.tag in ("TAVOLA", "TITOLO_TAVOLA"):
            assert (a.dxf.halign, a.dxf.valign) == (1, 2), f"{a.dxf.tag} non centrato"
            assert abs(a.dxf.align_point.x - ATTESI["centro_frontalino_x"]) < 0.01, a.dxf.align_point
    # linee guida per centrare i testi: layer non stampabile, disegnate per prime (dietro a tutto)
    guide = [i for i, e in enumerate(blk) if e.dxf.layer == "GUIDE"]
    assert len(guide) == ATTESI["guide"], f"linee guida trovate: {len(guide)}"
    assert guide == list(range(len(guide))), "le linee guida devono essere disegnate per prime"
    assert doc.layers.get("GUIDE").dxf.plot == 0, "il layer GUIDE non deve stampare"
    assert all(blk[i].dxf.get("color", 256) == 256 for i in guide)
    # cartiglio delle tavole: solo "cartiglio kevin", retino grigio per primo, campi come attributi
    cb = list(doc.blocks.get("PEMAJ_CARTIGLIO"))
    assert sorted(a.dxf.tag for a in cb if a.dxftype() == "ATTDEF") == sorted(ATTESI["attributi_cartiglio"])
    gr = [i for i, e in enumerate(cb) if genera_template.e_sfondo_grigio(e, doc)]
    assert gr and gr == list(range(len(gr))), "nel cartiglio il retino grigio deve stare dietro a tutto"
    assert r["cartiglio_mm"] == ATTESI["cartiglio_mm"], r


@test("TC5 A4 = intestazione; A3 = solo cartiglio; A2..A0 e LIBERO = stesso frontalino dell'A4")
def _():
    doc = ezdxf.readfile(str(tmp / "tpl.dxf"))
    nomi = [l.name for l in doc.layouts if l.name != "Model"]
    assert nomi == list(ATTESI["intestazione"]) + list(ATTESI["formati"]), nomi
    (n4, (W4, H4)), = ATTESI["intestazione"].items()
    lay = doc.layouts.get(n4)
    ins = [i.dxf.name for i in lay.query("INSERT")]
    assert ins == ["PEMAJ_CARTIGLIO_BASE"], ins
    assert tuple(round(v, 2) for v in lay.query("INSERT")[0].dxf.insert)[:2] == (0, H4)
    assert lay.dxf_layout.dxf.paper_width == W4 and lay.dxf_layout.dxf.paper_height == H4
    m = 3.46
    for nome, (W, H) in ATTESI["formati"].items():
        lay = doc.layouts.get(nome)
        ins = list(lay.query("INSERT"))
        if nome in ATTESI["solo_cartiglio"]:
            atteso, punto = "PEMAJ_CARTIGLIO", (round(W - m, 2), m)
        else:
            atteso, punto = "PEMAJ_CARTIGLIO_BASE", (0, H)
        assert [i.dxf.name for i in ins] == [atteso], (nome, [i.dxf.name for i in ins])
        assert tuple(round(v, 2) for v in ins[0].dxf.insert)[:2] == punto, (nome, ins[0].dxf.insert)
        sq = next(p for p in lay.query("LWPOLYLINE") if p.dxf.layer == "Squadratura")
        assert round(min(p[0] for p in sq.get_points()), 2) == m
        assert lay.dxf_layout.dxf.current_style_sheet == "PEMAJ_PROGETTI.ctb"
        vp = [v for v in lay.query("VIEWPORT") if v.dxf.id > 1]
        assert vp and round(vp[0].dxf.view_height / vp[0].dxf.height) == 100, nome
    assert doc.layouts.active_layout().name == n4, "il file deve aprirsi sull'intestazione A4"
    assert doc.header.get("$TILEMODE") == 0


@test("TC6 cartiglio con un altro committente: campi riconosciuti lo stesso")
def _():
    doc = ezdxf.readfile(str(CARTIGLIO))
    for t in doc.modelspace().query("TEXT"):
        if t.dxf.layer.lower() == "testo variabile":
            t.dxf.text = "ALTRO " + t.dxf.text
    doc.saveas(str(tmp / "cart2.dxf"))
    genera_template.genera(tmp / "cart2.dxf", tmp / "tpl2.dxf")
    d2 = ezdxf.readfile(str(tmp / "tpl2.dxf"))
    tags = {a.dxf.tag for a in d2.blocks.get("PEMAJ_CARTIGLIO_BASE").query("ATTDEF")}
    assert {"COMMITTENTE_4", "OGGETTO_3", "FASE"} <= tags, tags


@test("TC7 controllo tavola E01: risultati attesi")
def _():
    trov = verifica_tavola.verifica(T / "fixtures" / "ALLEGATO_E01.dxf")
    per = {}
    for t in trov:
        per[t[0]] = per.get(t[0], 0) + 1
    assert per == ATTESI["e01_per_codice"], per


@test("TC8 controllo tavola: errori costruiti apposta vengono trovati")
def _():
    doc = ezdxf.new("R2018")
    msp = doc.modelspace()
    doc.layers.add("FM", color=240)
    msp.add_text("CAPOVOLTO", height=25, rotation=180, dxfattribs={"layer": "FM"})
    msp.add_text("PICCOLO", height=10, dxfattribs={"layer": "FM"})
    msp.add_line((0, 0), (100, 0), dxfattribs={"layer": "FM", "color": 1})
    msp.add_line((0, 0), (0, 100), dxfattribs={"layer": "STRANO"})
    doc.saveas(str(tmp / "sbagliata.dxf"))
    codici = {t[0] for t in verifica_tavola.verifica(tmp / "sbagliata.dxf", "cm", 100, "diri")}
    assert {"V01", "V02", "V04", "V05", "V07", "V09"} <= codici, codici


@test("TC9 autoriparazione: script corrotto, CTB cancellati, cartiglio sparito")
def _():
    copia = tmp / "ws" / "tools" / "standard-pemaj"
    shutil.copytree(BASE, copia, ignore=shutil.ignore_patterns("out", "__pycache__"))
    man = copia / "manifest.json"
    if not man.exists():
        subprocess.run([sys.executable, str(copia / "dottore.py"), "--ripara"], capture_output=True)
        # impronte di prova per simulare un LOCK gia' congelato
        from comune import SCRIPT, sha256
        dati = {"script": {s: sha256(copia / s) for s in SCRIPT}}
        man.write_text(json.dumps(dati), encoding="utf-8")
        for s in SCRIPT:
            (copia / "_backup" / s).parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(copia / s, copia / "_backup" / s)
    # danni
    (copia / "genera_ctb.py").write_text("ROTTO", encoding="utf-8")
    for f in (copia / "out").glob("PEMAJ_*"):
        f.unlink()
    (copia / "riferimenti" / "CARTIGLIO BASE.dxf").unlink()
    alt = tmp / "ws" / "tools" / "planimetria-impianto" / "riferimenti"
    alt.mkdir(parents=True)
    shutil.copy2(CARTIGLIO, alt / "CARTIGLIO BASE.dxf")
    r = subprocess.run([sys.executable, str(copia / "dottore.py"), "--ripara"], capture_output=True, text=True, encoding="utf-8", errors="replace")
    esito = json.loads(r.stdout)
    assert esito["status"] == "riparato", esito
    testo = " | ".join(esito["azioni"])
    for atteso in ("ripristinato genera_ctb.py", "cartiglio recuperato", "rigenerati", "rigenerato"):
        assert atteso in testo, f"azione mancante: {atteso} -> {testo}"
    assert (copia / "genera_ctb.py").read_text(encoding="utf-8") != "ROTTO"
    assert (copia / "genera_ctb.py.modificato").exists(), "la versione danneggiata va conservata"
    r2 = subprocess.run([sys.executable, str(copia / "dottore.py")], capture_output=True, text=True, encoding="utf-8", errors="replace")
    assert json.loads(r2.stdout)["status"] == "ok", r2.stdout


@test("TC10 tabelle di base: cavo di terra giallo-verde, stile di quota, comando PEMAJ_FOGLIO")
def _():
    doc = ezdxf.readfile(str(tmp / "tpl.dxf"))
    ms = doc.mline_styles.get("PE_GIALLO_VERDE")
    assert [[e.offset, e.color, e.linetype] for e in ms.elements] == ATTESI["mline_pe"]
    assert "PE_TRATTO" in doc.linetypes
    assert doc.layers.get("PE_VERDE").dxf.linetype == "PE_TRATTO"
    assert doc.layers.get("PE_GIALLO").dxf.color == 50 and doc.layers.get("PE_VERDE").dxf.color == 94
    q = doc.dimstyles.get("PEMAJ")
    assert q.dxf.dimtxt == 2.5 and q.dxf.dimtxsty == "PEMAJ_2.5"
    lsp = (BASE / "PEMAJ_FOGLIO.lsp").read_text(encoding="utf-8")
    assert "c:PEMAJ_FOGLIO" in lsp and lsp.count("(") == lsp.count(")"), "LISP con parentesi sbilanciate"
    assert "vla-" not in lsp, "niente funzioni ActiveX: non esistono in AutoCAD LT"


@test("TC11 base dell'architetto: piante tenute, layer e colori studio, quote/retini/numeri via, mm")
def _():
    A = ATTESI["base_grana"]
    out = tmp / "base"
    rep = prepara_base.prepara(T / "fixtures" / "BASE_GRANA.dxf", out)
    assert rep["unita"]["ingresso"] == A["unita"] and rep["unita"]["fattore"] == 1000, rep["unita"]
    tenute = {v["titolo"] for v in rep["viste"] if v["tieni"]}
    scartate = {v["titolo"] for v in rep["viste"] if not v["tieni"]}
    assert set(A["piante"]) <= tenute, tenute
    assert set(A["scartate"]) <= scartate, scartate
    doc = ezdxf.readfile(rep["file_pulito"])
    msp = doc.modelspace()
    assert doc.header["$INSUNITS"] == 4
    tipi = {e.dxftype() for e in msp}
    assert not tipi & {"DIMENSION", "HATCH", "SOLID", "POINT", "IMAGE", "WIPEOUT"}, tipi
    assert {e.dxf.layer for e in msp} <= set(A["layer_ammessi"]), {e.dxf.layer for e in msp}
    assert all(e.dxf.get("color", 256) == 256 and e.dxf.get("lineweight", -1) == -1 for e in msp)
    testi = [" ".join(prepara_base.testo_di(e).split()) for e in msp.query("TEXT MTEXT")]
    mancanti = [t for t in A["locali"] if not any(t in x for x in testi)]
    assert not mancanti, f"nomi dei locali persi: {mancanti}"
    assert not any(x.replace(" ", "").isdigit() for x in testi), "numeri delle finestre rimasti"
    assert "MURI" in {e.dxf.layer for e in msp} and "SCRITTE" in {e.dxf.layer for e in msp}
    xs = [e.dxf.start.x for e in msp.query("LINE")]
    assert min(xs) > 2_000_000, "coordinate non convertite da m a mm"


@test("TC12 tabella dei layer corretta dall'utente: applicata e memorizzata per la volta dopo")
def _():
    from openpyxl import load_workbook
    vecchia = prepara_base.MAPPA_STUDIO
    prepara_base.MAPPA_STUDIO = tmp / "mappa_layer.json"
    try:
        out = tmp / "base2"
        rep = prepara_base.prepara(T / "fixtures" / "BASE_GRANA.dxf", out, solo_analisi=True)
        wb = load_workbook(rep["tabella_layer"])
        ws = wb["Layer"]
        for r in ws.iter_rows(min_row=2):
            if r[0].value == "proiez lontane":
                r[5].value = "ELIMINA"
        wb.save(out / "tabella.xlsx")
        rep2 = prepara_base.prepara(T / "fixtures" / "BASE_GRANA.dxf", out, tabella=out / "tabella.xlsx")
        assert json.loads(prepara_base.MAPPA_STUDIO.read_text(encoding="utf-8"))["proiez lontane"] == "ELIMINA"
        assert rep2["layer_dopo"].get("INF", 0) < 50, rep2["layer_dopo"]
        rep3 = prepara_base.prepara(T / "fixtures" / "BASE_GRANA.dxf", out / "terza")
        assert rep3["layer_dopo"] == rep2["layer_dopo"], "la mappa memorizzata non e' stata riusata"
    finally:
        prepara_base.MAPPA_STUDIO = vecchia


@test("TC13 macchinari inseriti sul layer dei muri -> MACCHINE; layer MISTO: linee lunghe muri, resto grigio")
def _():
    vecchia = prepara_base.MAPPA_STUDIO
    prepara_base.MAPPA_STUDIO = tmp / "mappa13.json"
    try:
        doc = ezdxf.new("R2018")
        doc.header["$INSUNITS"] = 4
        msp = doc.modelspace()
        mac = doc.blocks.new("POMPA_X")
        for k in range(200):
            mac.add_line((k, 0), (k, 50))
        porta = doc.blocks.new("PORTA_90")
        porta.add_line((0, 0), (900, 0))
        porta.add_arc((0, 0), 900, 0, 90)
        for x in (0, 20000):
            msp.add_lwpolyline([(x, 0), (x + 15000, 0), (x + 15000, 300), (x, 300)], close=True,
                               dxfattribs={"layer": "MURATURA"})
        msp.add_blockref("POMPA_X", (5000, 5000), dxfattribs={"layer": "MURO"})
        msp.add_blockref("PORTA_90", (2000, 0), dxfattribs={"layer": "PORTE"})
        msp.add_line((0, 8000), (6000, 8000), dxfattribs={"layer": "MIX"})       # 6 m: muro
        msp.add_line((0, 9000), (500, 9000), dxfattribs={"layer": "MIX"})        # 0,5 m: resto
        doc.saveas(str(tmp / "mix.dxf"))
        out = tmp / "mix"
        rep = prepara_base.prepara(tmp / "mix.dxf", out, solo_analisi=True)
        from openpyxl import load_workbook
        wb = load_workbook(rep["tabella_layer"])
        for r in wb["Layer"].iter_rows(min_row=2):
            if r[0].value == "MIX":
                r[5].value = "MISTO"
        blocchi = {r[0].value: r[4].value for r in wb["Blocchi"].iter_rows(min_row=2)}
        assert blocchi.get("POMPA_X") == "MACCHINE" and blocchi.get("PORTA_90") == "INF", blocchi
        wb.save(out / "t.xlsx")
        rep = prepara_base.prepara(tmp / "mix.dxf", out, tabella=out / "t.xlsx")
        d = ezdxf.readfile(rep["file_pulito"])
        ins = {i.dxf.name: i.dxf.layer for i in d.modelspace().query("INSERT")}
        assert ins == {"POMPA_X": "MACCHINE", "PORTA_90": "INF"}, ins
        lin = sorted((round((l.dxf.end - l.dxf.start).magnitude), l.dxf.layer) for l in d.modelspace().query("LINE"))
        assert lin == [(500, "INF"), (6000, "MURI")], lin
    finally:
        prepara_base.MAPPA_STUDIO = vecchia


ok = sum(1 for _, p, _ in risultati if p)
for nome, p, msg in risultati:
    print(("PASS " if p else "FAIL ") + nome + ("" if p else f"  -> {msg}"))
print(f"LOCK {ok}/{len(risultati)} {'PASS' if ok == len(risultati) else 'FAIL'}")
shutil.rmtree(tmp, ignore_errors=True)
sys.exit(0 if ok == len(risultati) else 1)
